from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from src import config, engine, loader, materials
from webapp import app as app_module
from webapp.storage import Store


ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "samples" / "样例企业-审计材料.xlsx"


def human_sheet(workbook, rows):
    if config.SHEET_HUMAN in workbook.sheetnames:
        del workbook[config.SHEET_HUMAN]
    ws = workbook.create_sheet(config.SHEET_HUMAN)
    ws.append(config.COL_HUMAN)
    for row in rows:
        ws.append(row)


def as_bytes(workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def human_workbook(month="2026-05") -> bytes:
    workbook = Workbook()
    ws = workbook.active
    ws.title = config.SHEET_COMPANY
    ws.append(["项目", "内容"])
    for label, value in zip(config.COMPANY_FIELDS, ("人力验收企业", "HUMAN-QA-001", "批发业", "2026H1")):
        ws.append([label, value])
    human_sheet(workbook, [
        ["个税申报", "隐私姓名甲", "SECRET-ID-001", "2026-05", "已申报", 10000, "个税原始台账"],
        ["个税申报", "隐私姓名乙", "SECRET-ID-002", "2026-05", "已申报", 9000, "个税原始台账"],
        ["社保参保", "隐私姓名甲", "SECRET-ID-001", month, "在保", 1800, "社保原始台账"],
        ["社保参保", "隐私姓名丙", "SECRET-ID-003", month, "停保", 0, "社保原始台账"],
        ["公积金缴存", "隐私姓名甲", "SECRET-ID-001", "2026-05", "正常缴存", 1200, "公积金原始台账"],
    ])
    return as_bytes(workbook)


class HumanRecordParsingTests(unittest.TestCase):
    def test_monthly_dedup_status_amount_and_rule(self):
        workbook = load_workbook(SAMPLE)
        if config.SHEET_SUPPLEMENT in workbook.sheetnames:
            sheet = workbook[config.SHEET_SUPPLEMENT]
            for row in range(sheet.max_row, 1, -1):
                if sheet.cell(row, 1).value in {"人力.个税申报人数", "人力.社保参保人数"}:
                    sheet.delete_rows(row)
        human_sheet(workbook, [
            ["个税申报", "张三", "ID-001", "2026-05", "已申报", 10000, "自然人电子税务局"],
            ["个税申报", "李四", "ID-002", "2026-05", "已申报", 9000, "自然人电子税务局"],
            ["社保参保", "张三", "ID-001", "2026-05", "正常参保", 1800, "社保平台"],
            ["社保参保", "王五", "ID-003", "2026-05", "停保", 0, "社保平台"],
            ["公积金缴存", "张三", "ID-001", "2026-05", "正常缴存", 1200, "公积金中心"],
        ])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "human.xlsx"
            workbook.save(path)
            workbook.close()
            dataset = loader.load(path)
        self.assertEqual(dataset.get("人力.个税申报人数"), 2)
        self.assertEqual(dataset.get("人力.社保参保人数"), 1)
        self.assertEqual(dataset.get("人力.公积金缴存人数"), 1)
        self.assertEqual(dataset.get("个税.工资薪金申报收入"), 19000)
        self.assertNotIn("ID-001", dataset.metrics["人力.个税申报人数"].detail)
        rule = next(rule for rule in engine.load_rules(ROOT / "rules") if rule.id == "R-004")
        self.assertEqual(engine.evaluate(rule, dataset).status, "hit")

    def test_latest_common_month_and_no_common_month_rejected(self):
        workbook = Workbook()
        human_sheet(workbook, [
            ["个税申报", "甲", "A-001", "2026-01", "已申报", 1, "个税"],
            ["社保参保", "甲", "A-001", "2026-01", "在保", 1, "社保"],
            ["个税申报", "甲", "A-001", "2026-02", "已申报", 1, "个税"],
            ["社保参保", "甲", "A-001", "2026-03", "在保", 1, "社保"],
        ])
        records = loader._read_human_records(workbook)
        company = loader.Company("测试", "T", "测试", "2026H1")
        metrics = loader._human_metrics(company, records)
        self.assertIn("核对月 2026-01", metrics["人力.个税申报人数"].detail)
        records = [record for record in records if record.month != "2026-01"]
        with self.assertRaisesRegex(loader.InputError, "没有相同所属月"):
            loader._human_metrics(company, records)
        workbook.close()

    def test_duplicate_negative_and_out_of_period_are_rejected(self):
        for rows, message in (
            ([
                ["社保参保", "甲", "A-001", "2026-05", "在保", 1, "社保"],
                ["社保参保", "甲", "A-001", "2026-05", "在保", 1, "社保"],
            ], "重复"),
            ([["社保参保", "甲", "A-001", "2026-05", "在保", -1, "社保"]], "不能为负数"),
        ):
            workbook = Workbook()
            human_sheet(workbook, rows)
            with self.assertRaisesRegex(loader.InputError, message):
                loader._read_human_records(workbook)
            workbook.close()
        record = loader._HumanRecord("社保", "a" * 64, "2025-12", True, None, "测试")
        with self.assertRaisesRegex(loader.InputError, "不在核对期间"):
            loader._human_metrics(loader.Company("测试", "T", "测试", "2026H1"), [record])

    def test_multifile_preview_keeps_only_hashed_identity_and_merges(self):
        source = Workbook()
        human_sheet(source, [
            ["个税申报", "张三", "SECRET-ID-001", "2026-05", "已申报", 10000, "个税系统"],
            ["社保参保", "张三", "SECRET-ID-001", "2026-05", "在保", 1800, "社保平台"],
        ])
        human_bytes = as_bytes(source)
        sample_bytes = SAMPLE.read_bytes()
        docs = materials.preview([("sample.xlsx", sample_bytes), ("human.xlsx", human_bytes)], set())
        self.assertFalse(any(doc["error"] for doc in docs), docs)
        serialized = str(docs[1]["human_records"])
        self.assertNotIn("SECRET-ID-001", serialized)
        self.assertNotIn("张三", serialized)
        selections = {doc["id"]: {"company": {}, "rows": doc["rows"]} for doc in docs}
        dataset = materials.build_dataset(docs, selections, {}, set())
        self.assertEqual(dataset.get("人力.个税申报人数"), 1)
        self.assertEqual(dataset.get("人力.社保参保人数"), 1)


class HumanRecordWebFlow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp.name) / "human.db"
        self.old_store = app_module.store
        app_module.store = Store(self.db_path)
        app_module.store.create_user("human-admin", "Human-test-2026!", "人力测试", "org_admin", "default")
        # A separate upload router owns its own bounded in-memory drafts. DB
        # replacement alone does not isolate the module-level app's closures.
        app = FastAPI(routes=[r for r in app_module.app.routes if not r.path.startswith("/api/materials/")],
                      exception_handlers=app_module.app.exception_handlers,
                      middleware=app_module.app.user_middleware)
        app_module.register_material_upload(app, app_module._user, app_module._allow,
                                            app_module._save_audit, app_module.RULES_DIR, app_module._audit_or_404)
        self.client = TestClient(app)
        self.client.post("/api/login", json={"username": "human-admin", "password": "Human-test-2026!"})

    def tearDown(self):
        self.client.close()
        app_module.store = self.old_store
        self.temp.cleanup()

    def audit(self, content):
        preview = self.client.post("/api/materials/preview", files=[("files", ("人力来源.xlsx", content))])
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertNotIn("隐私姓名", preview.text)
        self.assertNotIn("SECRET-ID", preview.text)
        draft = preview.json()
        doc = draft["documents"][0]
        return self.client.post("/api/materials/audit", json={
            "token": draft["token"], "mode": "merge", "same_scope": True,
            "company": doc["company"], "selections": [{"id": doc["id"], "company": doc["company"]}],
        })

    def test_human_evidence_privacy_archive_and_reopened_store(self):
        response = self.audit(human_workbook())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertFalse(response.json()["errors"])
        audit = response.json()["results"][0]["audit"]
        metric = {m["name"]: m for m in audit["metrics"]}
        self.assertEqual(metric["人力.个税申报人数"]["value"], "2.00")
        self.assertEqual(metric["人力.社保参保人数"]["value"], "1.00")
        self.assertEqual(metric["人力.公积金缴存人数"]["value"], "1.00")
        self.assertIn("核对月 2026-05", metric["人力.社保参保人数"]["detail"])
        self.assertIn("剔除 1 条", metric["人力.社保参保人数"]["detail"])
        finding = next(f for f in audit["findings"] if f["id"] == "R-004")
        self.assertEqual(finding["status"], "hit")
        self.assertIn("人力来源.xlsx", str(finding["evidence"]))
        audit_id = audit["audit_id"]
        html = self.client.get(f"/api/report/{audit_id}/html?version=1")
        self.assertEqual(html.status_code, 200)
        self.assertIn("人力来源.xlsx", html.text)
        app_module.store = Store(self.db_path)
        restored = self.client.get(f"/api/audits/{audit_id}")
        self.assertEqual(restored.json(), audit)
        self.assertEqual(self.client.get(f"/api/report/{audit_id}/html?version=1").content, html.content)
        saved = app_module.store.get_audit(audit_id)
        for output in (response.text, restored.text, html.text, str(saved["dataset"])):
            self.assertNotIn("隐私姓名", output)
            self.assertNotIn("SECRET-ID", output)
            self.assertNotIn(loader._human_person_key("SECRET-ID-001", "测试"), output)

    def test_no_common_month_does_not_create_audit(self):
        response = self.audit(human_workbook(month="2026-04"))
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["results"], [])
        self.assertIn("没有相同所属月", str(response.json()["errors"]))
        with app_module.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM audits").fetchone()[0], 0)


if __name__ == "__main__":
    unittest.main()
