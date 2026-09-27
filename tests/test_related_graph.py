"""FR-B09 graph traversal stays separate from YAML amount/ratio rules."""
from __future__ import annotations

import tempfile
import unittest
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src import config, engine, loader, materials, related_graph, render
from webapp import app as app_module
from webapp.knowledge import build_graph
from webapp.storage import Store, deserialize_dataset, serialize_dataset

ROOT = Path(__file__).resolve().parent.parent


def graph_workbook():
    wb = load_workbook(ROOT / "samples" / "样例企业-审计材料.xlsx")
    company = loader._read_company(wb)
    ws = wb.create_sheet(related_graph.SHEET_SUBJECTS)
    ws.append(related_graph.SUBJECT_HEADERS)
    ws.append(["A", company.name, "企业", company.taxpayer_id])
    ws.append(["P", "仿真股东甲", "个人", ""])
    ws.append(["B", "仿真关联公司乙", "企业", "SIM-OTHER-01"])
    ws = wb.create_sheet(related_graph.SHEET_RELATIONS)
    ws.append(related_graph.RELATION_HEADERS)
    ws.append(["S-1", "P", "A", "股东", "2025-01-01", "", "已复核", "仿真股权登记摘录"])
    ws.append(["C-1", "P", "B", "控制", "2025-01-01", "", "已复核", "仿真控制关系底稿"])
    ws = wb.create_sheet(related_graph.SHEET_TRADES)
    ws.append(related_graph.TRADE_HEADERS)
    ws.append(["T-1", "A", "B", "2026-03-20", 100000, "仿真可比价格偏离待核实", "已复核"])
    return wb


class RelatedGraphTests(unittest.TestCase):
    def _load(self, wb):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "graph.xlsx"
            wb.save(path)
            wb.close()
            return loader.load(path)

    def test_reviewed_shareholder_control_trade_path_hits_independent_rule(self):
        data = self._load(graph_workbook())
        findings = related_graph.run(data)
        self.assertEqual(len(findings), 1)
        self.assertEqual((findings[0].rule.id, findings[0].status), ("G-001", "hit"))
        self.assertIn("T-1", findings[0].evidence[-1].value)
        self.assertIn("关联交易!第2行", findings[0].evidence[-1].source)
        yaml_results = engine.run(engine.load_rules(ROOT / "rules"), data)
        self.assertEqual(len(yaml_results), 24)
        self.assertTrue(all(item.rule.id.startswith("R-") for item in yaml_results))

    def test_unreviewed_relation_or_unsubstantiated_trade_never_hits(self):
        wb = graph_workbook()
        wb[related_graph.SHEET_RELATIONS]["G3"] = "待复核"
        data = self._load(wb)
        self.assertEqual(related_graph.run(data)[0].status, "skipped")
        wb = graph_workbook()
        wb[related_graph.SHEET_RELATIONS]["F3"] = "2026-01-31"
        data = self._load(wb)
        self.assertEqual(related_graph.run(data)[0].status, "skipped")
        wb = graph_workbook()
        wb[related_graph.SHEET_TRADES]["F2"] = ""
        data = self._load(wb)
        self.assertEqual(related_graph.run(data)[0].status, "skipped")

    def test_multiple_shareholding_periods_do_not_hide_valid_path(self):
        wb = graph_workbook()
        wb[related_graph.SHEET_RELATIONS].append(
            ["S-2", "P", "A", "股东", "2027-01-01", "", "已复核", "后续期间股权记录"]
        )
        data = self._load(wb)
        paths = list(related_graph.candidate_paths(data))
        self.assertEqual([share.key for share, _, _ in paths], ["S-1"])
        self.assertEqual(related_graph.run(data)[0].status, "hit")

    def test_many_paths_count_all_hits_but_bound_evidence(self):
        wb = graph_workbook()
        sheet = wb[related_graph.SHEET_TRADES]
        for number in range(2, 27):
            sheet.append([f"T-{number}", "A", "B", "2026-03-20", 100000,
                          "仿真可比价格偏离待核实", "已复核"])
        data = self._load(wb)
        finding = related_graph.run(data)[0]
        self.assertIn("26 条", finding.conclusion)
        self.assertEqual(len(finding.evidence), 60)

    def test_network_limit_is_explicit_without_truncating_frozen_rule_result(self):
        for count in (50, 51):
            with self.subTest(count=count):
                wb = graph_workbook()
                for number in range(2, count + 1):
                    wb[related_graph.SHEET_TRADES].append(
                        [f"T-{number}", "A", "B", "2026-03-20", 100000, "已复核异常依据", "已复核"]
                    )
                data = self._load(wb)
                findings = related_graph.run(data)
                network = build_graph([], {"id": "bounded", "dataset": data, "findings": findings})
                self.assertEqual(network["related_paths"], {"shown": 50, "limit": 50, "truncated": count > 50})
                self.assertEqual(len([n for n in network["nodes"] if n["kind"] == "trade"]), 50)
                result = next(n for n in network["nodes"] if n["id"] == "finding:G-001")
                self.assertIn(f"{count} 条", result["conclusion"])
                ids = {n["id"] for n in network["nodes"]}
                self.assertTrue(all(e["source"] in ids and e["target"] in ids for e in network["edges"]))

    def test_network_trade_date_direction_and_unreviewed_reason(self):
        wb = graph_workbook()
        ws = wb[related_graph.SHEET_TRADES]
        ws["B2"], ws["C2"], ws["G2"] = "B", "A", "待复核"
        data = self._load(wb)
        network = build_graph([], {"id": "pending", "dataset": data, "findings": related_graph.run(data)})
        trade = next(n for n in network["nodes"] if n["kind"] == "trade")
        self.assertEqual(trade["period"], "2026-03-20")
        self.assertEqual(trade["status"], "skipped")
        self.assertEqual(trade["reason"], "交易尚未复核")
        self.assertIn({"source": "related:trade:T-1", "target": "related:subject:B", "label": "销售方"}, network["edges"])
        self.assertFalse(any(e["label"] == "关联方图规则命中" for e in network["edges"]))

    def test_graph_http_scope_reopen_and_current_assignment(self):
        data = self._load(graph_workbook())
        with tempfile.TemporaryDirectory() as tmp:
            previous = app_module.store
            try:
                db = Path(tmp) / "graph-http.db"
                app_module.store = store = Store(db)
                admin = store.create_user("graphowner", "Graph-test-2026!", "本机构", "org_admin", "graph-a")
                accountant = store.create_user("graphstaff", "Graph-test-2026!", "会计", "accountant", "graph-a")
                successor = store.create_user("graphnext", "Graph-test-2026!", "新负责人", "accountant", "graph-a")
                foreign = store.create_user("graphforeign", "Graph-test-2026!", "其他机构", "org_admin", "graph-b")
                customer = store.upsert_client(admin, data.company.name, data.company.taxpayer_id, accountant["id"])
                result = app_module._save_audit(data, admin, customer["id"])
                audit_id = result["audit_id"]
                with TestClient(app_module.app) as client:
                    _, token = store.authenticate(accountant["username"], "Graph-test-2026!")
                    client.cookies.set(app_module.COOKIE_NAME, token)
                    url = f"/api/knowledge/graph?audit_id={audit_id}"
                    before = client.get(url)
                    self.assertEqual(before.status_code, 200, before.text)
                    self.assertIn("仿真股东甲", before.text)
                    app_module.store = Store(db)
                    self.assertEqual(client.get(url).json(), before.json())
                    store.upsert_client(admin, data.company.name, data.company.taxpayer_id, successor["id"])
                    self.assertEqual(client.get(url).status_code, 404)
                    _, foreign_token = store.authenticate(foreign["username"], "Graph-test-2026!")
                    client.cookies.clear()
                    client.cookies.set(app_module.COOKIE_NAME, foreign_token)
                    self.assertEqual(client.get(url).status_code, 404)
                    public = client.get("/api/knowledge/graph")
                    self.assertEqual(public.status_code, 200)
                    self.assertNotIn("仿真股东甲", public.text)
                    self.assertNotIn("related:trade:", public.text)
            finally:
                app_module.store = previous

    def test_bad_reference_and_out_of_period_trade_are_rejected(self):
        wb = graph_workbook()
        wb[related_graph.SHEET_RELATIONS]["C3"] = "MISSING"
        with self.assertRaisesRegex(loader.InputError, "主体引用无效"):
            self._load(wb)
        wb = graph_workbook()
        wb[related_graph.SHEET_TRADES]["D2"] = "2025-03-20"
        with self.assertRaisesRegex(loader.InputError, "不在审计所属期"):
            self._load(wb)
        wb = graph_workbook()
        wb[related_graph.SHEET_RELATIONS]["F3"] = "2024-01-01"
        with self.assertRaisesRegex(loader.InputError, "终止日早于起始日"):
            self._load(wb)

    def test_excel_date_cells_are_accepted_but_non_iso_strings_are_rejected(self):
        wb = graph_workbook()
        wb[related_graph.SHEET_RELATIONS]["E2"] = datetime(2025, 1, 1)
        wb[related_graph.SHEET_TRADES]["D2"] = datetime(2026, 3, 20)
        data = self._load(wb)
        self.assertEqual(related_graph.run(data)[0].status, "hit")
        wb = graph_workbook()
        wb[related_graph.SHEET_RELATIONS]["E2"] = "20250101"
        with self.assertRaisesRegex(loader.InputError, "YYYY-MM-DD"):
            self._load(wb)
        wb = graph_workbook()
        wb[related_graph.SHEET_TRADES]["D2"] = "20260320"
        with self.assertRaisesRegex(loader.InputError, "YYYY-MM-DD"):
            self._load(wb)

    def test_huge_trade_amount_is_rejected_before_report_formatting(self):
        wb = graph_workbook()
        wb[related_graph.SHEET_TRADES]["E2"] = "1e1000000"
        with self.assertRaisesRegex(loader.InputError, "交易金额须为正数"):
            self._load(wb)

    def test_graph_only_material_can_be_audited(self):
        wb = graph_workbook()
        keep = {config.SHEET_COMPANY, related_graph.SHEET_SUBJECTS,
                related_graph.SHEET_RELATIONS, related_graph.SHEET_TRADES}
        for name in list(wb.sheetnames):
            if name not in keep:
                del wb[name]
        stream = BytesIO()
        wb.save(stream)
        wb.close()
        keys = {key for rule in engine.load_rules(ROOT / "rules") for key in rule.inputs}
        docs = materials.preview([("graph-only.xlsx", stream.getvalue())], keys)
        self.assertFalse(docs[0]["error"])
        selections = {docs[0]["id"]: {"id": docs[0]["id"], "company": deepcopy(docs[0]["company"]),
                                      "rows": [], "reviewed": True}}
        data = materials.build_dataset(docs, selections, {}, keys)
        self.assertEqual(data.metrics, {})
        self.assertEqual(related_graph.run(data)[0].status, "hit")

    def test_partial_graph_and_formula_are_rejected(self):
        wb = graph_workbook()
        del wb[related_graph.SHEET_TRADES]
        with self.assertRaisesRegex(loader.InputError, "同时包含"):
            self._load(wb)
        wb = graph_workbook()
        wb[related_graph.SHEET_SUBJECTS]["B3"] = '=1+1'
        with self.assertRaisesRegex(loader.InputError, "含公式"):
            self._load(wb)

    def test_material_preview_merge_and_audit_snapshot_keep_graph_sources(self):
        wb = graph_workbook()
        stream = BytesIO()
        wb.save(stream)
        wb.close()
        rules = engine.load_rules(ROOT / "rules")
        keys = {key for rule in rules for key in rule.inputs}
        docs = materials.preview([("graph.xlsx", stream.getvalue())], keys)
        self.assertFalse(docs[0]["error"])
        selections = {doc["id"]: {"id": doc["id"], "company": deepcopy(doc["company"]),
                                  "rows": deepcopy(doc["rows"]), "reviewed": True} for doc in docs}
        data = materials.build_dataset(docs, selections, {}, keys)
        self.assertIn("SHA256", data.related_graph.trades[0].source)
        restored = deserialize_dataset(serialize_dataset(data))
        self.assertEqual(related_graph.run(restored)[0].status, "hit")
        with tempfile.TemporaryDirectory() as tmp:
            previous = app_module.store
            try:
                app_module.store = Store(Path(tmp) / "graph.db")
                user = app_module.store.create_user("graphadmin", "Graph-pass-2026!", "测试管理员", "org_admin", "org-graph")
                result = app_module._save_audit(data, user)
                self.assertTrue(any(item["id"] == "G-001" and item["status"] == "hit"
                                    for item in result["findings"]))
                saved = app_module.store.get_audit(result["audit_id"])
                self.assertEqual(related_graph.run(saved["dataset"])[0].status, "hit")
                network = build_graph([], saved)
                self.assertTrue({"entity", "relation", "trade"}.issubset({node["kind"] for node in network["nodes"]}))
                self.assertIn({"source": "related:trade:T-1", "target": "finding:G-001",
                               "label": "关联方图规则命中"}, network["edges"])
                html, _ = render.render_html(saved["dataset"], saved["findings"], write=False)
                self.assertIn("G-001", html)
                self.assertIn("仿真股东甲", html)
            finally:
                app_module.store = previous


if __name__ == "__main__":
    unittest.main()
