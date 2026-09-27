from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from fastapi.testclient import TestClient
from openpyxl import Workbook

from src import engine, materials
from webapp import app as app_module
from webapp.storage import Store


ROOT = Path(__file__).resolve().parent.parent
COMPANY = {"name": "发票适配测试企业", "taxpayer_id": "BUYER-TEST-001", "industry": "批发业", "period": "2026-01"}
KEYS = {key for rule in engine.load_rules(ROOT / "rules") for key in rule.inputs}


def selections(docs):
    return {
        doc["id"]: {
            "id": doc["id"],
            "company": deepcopy(doc["company"]),
            "rows": deepcopy(doc["rows"]),
            "reviewed": True,
        }
        for doc in docs
    }


def invoice_workbook(rows, headers=None, sheet="电子税务局发票查询", company=COMPANY):
    headers = headers or [
        "数电票号码", "开票时间", "发票方向", "合计金额", "合计税额", "价税合计",
        "发票状态", "用途确认状态", "用途确认", "用途确认所属期", "是否进项税额转出", "进项转出税额",
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    if company:
        ws = workbook.create_sheet("企业信息")
        ws.append(["项目", "内容"])
        for label, key in zip(("企业名称", "纳税人识别号", "所属行业", "所属期"), company):
            ws.append([label, company[key]])
    ws = workbook.create_sheet(sheet)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def invoice_xml(
    number="24442000000000000001",
    amount="100.00",
    tax="13.00",
    total="113.00",
    issued="2026-01-10T09:30:00",
    red="false",
    confirmed="true",
    usage="抵扣税款",
    period="2026-01",
    transferred="false",
    transferred_tax="0",
):
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<xbrli:xbrl xmlns:xbrli="http://www.xbrl.org/2003/instance" xmlns:einv="http://xbrl.mof.gov.cn/taxonomy/2023-12-31/einv">
  <einv:InvoiceNumber>{number}</einv:InvoiceNumber>
  <einv:SellerName>测试供应商</einv:SellerName>
  <einv:SellerIdNum>SELLER-TEST-001</einv:SellerIdNum>
  <einv:UnifiedSocialCreditCodeOfAccountingEntity>{COMPANY["taxpayer_id"]}</einv:UnifiedSocialCreditCodeOfAccountingEntity>
  <einv:NameOfAccountingEntity>{COMPANY["name"]}</einv:NameOfAccountingEntity>
  <einv:RequestTime>{issued}</einv:RequestTime>
  <einv:TotalAmWithoutTax>{amount}</einv:TotalAmWithoutTax>
  <einv:TotalTaxAm>{tax}</einv:TotalTaxAm>
  <einv:TotalTax-includedAmount>{total}</einv:TotalTax-includedAmount>
  <einv:WhetherEinvoiceIsRedEinvoice>{red}</einv:WhetherEinvoiceIsRedEinvoice>
  <einv:WhetherEinvoiceUsageHasBeenConfirmed>{confirmed}</einv:WhetherEinvoiceUsageHasBeenConfirmed>
  <einv:UsageConfirmation>{usage}</einv:UsageConfirmation>
  <einv:PeriodOfUsageConfirmation>{period}</einv:PeriodOfUsageConfirmation>
  <einv:WhetherInputVatHasBeenTransferredOut>{transferred}</einv:WhetherInputVatHasBeenTransferredOut>
  <einv:AmountOfTransferredOutInputVat>{transferred_tax}</einv:AmountOfTransferredOutInputVat>
</xbrli:xbrl>'''.encode()


class InvoiceSourceParsing(unittest.TestCase):
    def test_source_excel_maps_red_void_and_tax_basis(self):
        rows = [
            ["S-1", "2026-01-02", "销项", 100, 13, 113, "正常", "", "", "", "", ""],
            ["S-2", "2026-01-03", "销项", 20, 2.6, 22.6, "红字", "", "", "", "", ""],
            ["S-3", "2026-01-04", "销项", 50, 6.5, 56.5, "作废", "", "", "", "", ""],
            ["P-1", "2026-01-05", "采购", 60, 7.8, 67.8, "正常", "已确认", "抵扣税款", "2026-01", "否", 0],
            ["P-2", "2026-01-06", "采购", 40, 5.2, 45.2, "正常", "未确认", "", "", "否", 0],
            ["P-3", "2026-01-07", "采购", 10, 1.3, 11.3, "正常", "已确认", "抵扣税款", "2026-01", "是", 0.3],
        ]
        doc = materials.preview([("电子税务局导出.xlsx", invoice_workbook(rows))], KEYS)[0]
        self.assertFalse(doc["error"], doc)
        self.assertEqual(len(doc["invoices"]), 6)
        data = materials.build_dataset([doc], selections([doc]), {}, KEYS)
        self.assertEqual(data.get("发票.销项净额"), 80)
        self.assertEqual(data.get("发票.销项税额净额"), Decimal("10.4"))
        self.assertEqual(data.get("发票.采购不含税净额"), 110)
        self.assertEqual(data.get("发票.采购税额净额"), Decimal("14.3"))
        self.assertEqual(data.get("凭证.本期确认抵扣税额"), Decimal("8.8"))
        self.assertIn("剔除作废 1 张", data.detail_of("发票.销项净额"))
        self.assertIn("未确认用途的采购税额不计入", data.detail_of("凭证.本期确认抵扣税额"))
        findings = engine.run(engine.load_rules(ROOT / "rules"), data)
        self.assertEqual(next(finding for finding in findings if finding.rule.id == "R-019").status, "pass")

    def test_excel_rejects_duplicate_number_and_tax_inclusive_mismatch(self):
        duplicate = [
            ["DUP-1", "2026-01-02", "销项", 100, 13, 113, "正常"],
            ["DUP-1", "2026-01-03", "销项", 20, 2.6, 22.6, "红字"],
        ]
        doc = materials.preview([("重复.xlsx", invoice_workbook(duplicate))], KEYS)[0]
        self.assertIn("发票号码重复", doc["error"])
        mismatch = [["BAD-1", "2026-01-02", "销项", 100, 13, 999, "正常"]]
        doc = materials.preview([("税额口径.xlsx", invoice_workbook(mismatch))], KEYS)[0]
        self.assertIn("价税合计", doc["error"])

    def test_mof_xbrl_xml_merge_red_and_duplicate_detection(self):
        docs = materials.preview([
            ("采购蓝票.xml", invoice_xml()),
            ("采购红票.xml", invoice_xml(
                number="24442000000000000002", amount="20", tax="2.6", total="22.6", red="true"
            )),
        ], KEYS)
        self.assertTrue(all(not doc["error"] for doc in docs), docs)
        data = materials.build_dataset(docs, selections(docs), COMPANY, KEYS)
        self.assertEqual(data.get("发票.采购不含税净额"), 80)
        self.assertEqual(data.get("发票.采购税额净额"), Decimal("10.4"))
        self.assertEqual(data.get("凭证.本期确认抵扣税额"), Decimal("10.4"))
        self.assertIn("采购蓝票.xml", data.source_of("发票.采购不含税净额"))
        self.assertIn("采购红票.xml", data.source_of("发票.采购不含税净额"))

        duplicate = materials.preview([("a.xml", invoice_xml()), ("b.xml", invoice_xml())], KEYS)
        with self.assertRaisesRegex(materials.InputError, "跨文件重复"):
            materials.build_dataset(duplicate, selections(duplicate), COMPANY, KEYS)

    def test_xml_unconfirmed_tax_is_not_deductible_and_period_is_enforced(self):
        doc = materials.preview([("未确认.xml", invoice_xml(confirmed="false", usage=""))], KEYS)[0]
        data = materials.build_dataset([doc], selections([doc]), COMPANY, KEYS)
        self.assertEqual(data.get("发票.采购税额净额"), 13)
        self.assertIsNone(data.get("凭证.本期确认抵扣税额"))

        outside = materials.preview([("跨期.xml", invoice_xml(issued="2025-12-31"))], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "不在核对期间"):
            materials.build_dataset([outside], selections([outside]), COMPANY, KEYS)

    def test_xml_security_and_required_fields(self):
        entity = b'''<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e "boom">]><x><InvoiceNumber>&e;</InvoiceNumber></x>'''
        self.assertIn("DTD", materials.preview([("实体.xml", entity)], KEYS)[0]["error"])
        fake = b'''<xbrl xmlns:einv="https://example.test/einv"><einv:InvoiceNumber>1</einv:InvoiceNumber></xbrl>'''
        self.assertIn("财政部", materials.preview([("非标准.xml", fake)], KEYS)[0]["error"])
        mismatch = invoice_xml(total="999")
        self.assertIn("价税合计", materials.preview([("错误.xml", mismatch)], KEYS)[0]["error"])


class InvoiceWebFlow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_store = app_module.store
        app_module.store = Store(Path(self.temp.name) / "invoice.db")
        app_module.store.create_user("invoice-admin", "invoice-test-2026", "发票测试", "org_admin", "default")
        self.client = TestClient(app_module.app)
        self.client.post("/api/login", json={"username": "invoice-admin", "password": "invoice-test-2026"})

    def tearDown(self):
        self.client.close()
        app_module.store = self.old_store
        self.temp.cleanup()

    def test_xml_and_source_excel_commit_with_persisted_evidence(self):
        sales = invoice_workbook([
            ["S-WEB", "2026-01-02", "销项", 200, 26, 226, "正常", "", "", "", "", ""],
        ])
        response = self.client.post(
            "/api/materials/preview",
            files=[("files", ("销项导出.xlsx", sales)), ("files", ("采购.xml", invoice_xml()))],
        )
        self.assertEqual(response.status_code, 200, response.text)
        draft = response.json()
        self.assertEqual([doc["kind"] for doc in draft["documents"]], ["xlsx", "xml"])
        payload = {
            "token": draft["token"], "mode": "merge", "same_scope": True,
            "company": COMPANY, "selections": list(selections(draft["documents"]).values()),
        }
        audited = self.client.post("/api/materials/audit", json=payload)
        self.assertEqual(audited.status_code, 200, audited.text)
        body = audited.json()
        self.assertFalse(body["errors"], body)
        audit_id = body["results"][0]["audit"]["audit_id"]
        saved = app_module.store.get_audit(audit_id)["dataset"]
        self.assertEqual(saved.get("发票.销项净额"), 200)
        self.assertEqual(saved.get("发票.采购不含税净额"), 100)
        self.assertIn("采购.xml", saved.source_of("发票.采购不含税净额"))


if __name__ == "__main__":
    unittest.main()
