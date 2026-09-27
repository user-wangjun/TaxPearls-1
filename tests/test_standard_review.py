"""Standard cell review preserves original bytes and deterministic derivation."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import load_workbook

from src import config, engine, materials, material_review
from src.input_errors import InputError
from src.snapshots import deserialize_dataset
from tests.test_materials import accounts, workbook, COMPANY, KEYS
from webapp.enterprise_analysis import analyze, prefill


class StandardReviewTests(unittest.TestCase):
    def setUp(self):
        self.rules = engine.load_rules(Path(__file__).resolve().parents[1] / 'rules')

    def preview(self, raw):
        docs = materials.preview([('账表.xlsx', raw)], KEYS, allow_incomplete_company=True, capture_standard=True)
        self.assertFalse(docs[0]['error'], docs[0]['error'])
        return docs

    def review(self, docs, corrections):
        return analyze(docs, prefill(docs), {'0': {'standard_edits': corrections}}, self.rules)

    def test_account_cell_recalculates_and_keeps_original_value_and_coordinates(self):
        docs = self.preview(accounts())
        original = deepcopy(docs)
        result = self.review(docs, {'科目余额表!E2': '80000'})
        self.assertTrue(result['can_confirm'], result['feedback'])
        data = deserialize_dataset(result['dataset'])
        self.assertEqual(data.get('营业收入'), 80000)
        self.assertIn('科目余额表!E2', data.source_of('营业收入'))
        self.assertIn('100000', data.source_of('营业收入'))
        self.assertIn('用户修正', data.source_of('营业收入'))
        change = next(e for e in result['edits'] if e['field'] == '科目余额表!E2')
        self.assertEqual((change['old'], change['new']), ('100000', '80000'))
        self.assertEqual(docs, original)

    def test_blank_declaration_can_be_filled_zero_is_not_blank(self):
        docs = self.preview(workbook('增值税申报', [['项目', '金额'], ['销售额', None]], COMPANY))
        self.assertFalse(analyze(docs, prefill(docs), {}, self.rules)['can_confirm'])
        result = self.review(docs, {'增值税申报!B2': '0'})
        self.assertTrue(result['can_confirm'])
        data = deserialize_dataset(result['dataset'])
        self.assertEqual(data.get('增值税.销售额'), 0)
        self.assertIn('原文空白', data.source_of('增值税.销售额'))
        self.assertIn('用户补填', data.source_of('增值税.销售额'))
        self.assertFalse(self.review(docs, {'增值税申报!B2': ''})['can_confirm'])

    def test_clearing_account_value_removes_derived_metric_not_zero(self):
        book = load_workbook(BytesIO(accounts()))
        tax = book.create_sheet('增值税申报')
        tax.append(['项目', '金额'])
        tax.append(['销售额', 100000])
        raw = BytesIO()
        book.save(raw)
        book.close()
        result = self.review(self.preview(raw.getvalue()), {'科目余额表!E2': ''})
        self.assertTrue(result['can_confirm'])
        self.assertIsNone(deserialize_dataset(result['dataset']).get('营业收入'))
        check = next(c for c in result['checks'] if c['rule_id'] == 'R-001')
        self.assertFalse(check['ready'])

    def test_statement_and_history_corrections_rederive_trend_without_current_contamination(self):
        raw = workbook('利润表', [config.COL_STATEMENT, ['营业收入', 100000]], COMPANY)
        history = workbook('利润表', [config.COL_STATEMENT, ['营业收入', 300000]], {**COMPANY, 'period': '2025-01'})
        docs = materials.preview([('本期.xlsx', raw), ('历史.xlsx', history)], KEYS, capture_standard=True)
        result = analyze(docs, prefill([docs[0]]), {
            '0': {'standard_edits': {'利润表!B2': '120000'}},
            '1': {'purpose': 'history', 'standard_edits': {'利润表!B2': '240000'}}}, self.rules)
        self.assertTrue(result['can_confirm'], result['feedback'])
        data = deserialize_dataset(result['dataset'])
        self.assertEqual(data.get('利润表.营业收入'), 120000)
        self.assertEqual(data.get('趋势.营业收入.上年同期'), 240000)
        self.assertIn('用户修正', data.source_of('趋势.营业收入.上年同期'))
        self.assertIn('历史.xlsx', data.source_of('趋势.营业收入.上年同期'))

    def test_period_series_edit_and_clear_recalculate_instead_of_leaving_old_derived_values(self):
        book = load_workbook(BytesIO(workbook('利润表', [config.COL_STATEMENT, ['营业收入', 100]], COMPANY)))
        sheet = book.create_sheet('历史指标')
        sheet.append(config.COL_HISTORY)
        sheet.append(['利润表.营业收入', 200, '原报表', '2025-01', '同口径'])
        raw = BytesIO()
        book.save(raw)
        book.close()
        docs = self.preview(raw.getvalue())
        for value, expected in [('250', 250), ('', None)]:
            result = self.review(docs, {'历史指标!B2': value})
            self.assertTrue(result['can_confirm'], result['feedback'])
            self.assertEqual(deserialize_dataset(result['dataset']).get('趋势.营业收入.上年同期'), expected)

    def test_supplement_blank_requires_original_source_and_unit_notes_on_fill(self):
        for source, detail, allowed in [('来源底稿', '元，同口径', True), ('', '元', False), ('来源底稿', '', False)]:
            docs = self.preview(workbook('补充指标', [config.COL_SUPPLEMENT,
                ['人力.社保参保人数', None, source, COMPANY['period'], detail]], COMPANY))
            result = self.review(docs, {'补充指标!B2': '3'})
            self.assertEqual(result['can_confirm'], allowed, result['feedback'])
            if allowed:
                self.assertEqual(deserialize_dataset(result['dataset']).get('人力.社保参保人数'), 3)
                self.assertFalse(self.review(docs, {'补充指标!B2': '1.5'})['can_confirm'])

    def test_untrusted_fields_formula_boolean_nonfinite_and_oversize_rejected(self):
        docs = self.preview(accounts())
        for corrections in [[], {'科目余额表!A2': '9999'}, {'不存在!B2': '0'},
                            {'科目余额表!E2': True}, {'科目余额表!E2': {'source': 'forged'}},
                            {'科目余额表!E2': '=1+1'}, {'科目余额表!E2': 'NaN'},
                            {'科目余额表!E2': '1e100000'}, {'科目余额表!E2': '9' * 101}]:
            with self.subTest(corrections=corrections), self.assertRaises(InputError):
                self.review(docs, corrections)

    def test_source_annotation_matches_whole_coordinates_only(self):
        tables = {'利润表': [config.COL_STATEMENT, ['营业收入', 100]]}
        document = {'rows': [{'source': '利润表!B2 与 利润表!B20', 'detail': '合并来源'}]}
        material_review.annotate(document, tables, {'利润表!B2': '80'})
        self.assertEqual(document['rows'][0]['source'].count('用户修正'), 1)
        self.assertIn('与 利润表!B20', document['rows'][0]['source'])

    def test_corrections_do_not_hide_conflicting_material_or_change_other_field_source(self):
        docs = materials.preview([('A.xlsx', accounts()), ('B.xlsx', accounts(COMPANY, 80000))], KEYS, capture_standard=True)
        result = analyze(docs, prefill(docs), {'0': {'standard_edits': {'科目余额表!E2': '90000'}}}, self.rules)
        self.assertFalse(result['can_confirm'])
        self.assertIn('input_conflict', [n['code'] for n in result['feedback']['blocking']])
        result = self.review([docs[0]], {'科目余额表!C2': '100'})
        self.assertTrue(result['can_confirm'])
        self.assertEqual(deserialize_dataset(result['dataset']).get('营业收入'), 100000)
        self.assertNotIn('用户修正', deserialize_dataset(result['dataset']).source_of('营业收入'))


if __name__ == '__main__':
    unittest.main()
