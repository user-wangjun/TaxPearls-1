from __future__ import annotations

import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src import engine, loader, training
from webapp import app as app_module
from webapp.storage import Store

ROOT = Path(__file__).resolve().parent.parent


class P1StorageAndAuth(unittest.TestCase):
    def test_argon2_session_and_org_scoping(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "p1.db"
            repo = Store(path)
            admin = repo.create_user("admin", "correct-horse-2026", "管理员", "org_admin", "org-a")
            accountant = repo.create_user("accountant", "accountant-2026", "会计", "accountant", "org-a")
            outsider = repo.create_user("outsider", "outsider-pass-2026", "外部会计", "accountant", "org-b")
            self.assertIsNone(repo.authenticate("admin", "wrong-password"))
            auth = repo.authenticate("admin", "correct-horse-2026")
            self.assertIsNotNone(auth)
            user, token = auth
            self.assertEqual(repo.user_for_token(token)["id"], admin["id"])
            client = repo.upsert_client(admin, "测试客户", "TEST-001", accountant["id"])
            self.assertEqual([c["id"] for c in repo.list_clients(accountant)], [client["id"]])
            self.assertEqual(repo.list_clients(outsider), [])
            with closing(sqlite3.connect(path)) as db:
                password_hash = db.execute("SELECT password_hash FROM users WHERE id=?", (admin["id"],)).fetchone()[0]
                stored_token = db.execute("SELECT token_hash FROM sessions").fetchone()[0]
            self.assertTrue(password_hash.startswith("$argon2id$"))
            self.assertNotEqual(stored_token, token)


class P1InputsAndTraining(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rules = engine.load_rules(ROOT / "rules")

    def test_structured_p1_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            wb = load_workbook(ROOT / "samples" / "样例企业-审计材料.xlsx")
            ws = wb.create_sheet("利润表")
            ws.append(["项目", "本期金额"])
            ws.append(["营业收入", 1_280_000])
            ws.append(["营业成本", 780_000])
            ws = wb.create_sheet("历史指标")
            ws.append(["指标", "数值", "来源", "所属期", "口径说明"])
            ws.append(["历史.上年同期收入", 1_000_000, "上年度利润表", "2025H1", "同口径半年数"])
            ws = wb.create_sheet("发票明细")
            ws.append(["发票号码", "开票日期", "类型", "不含税金额", "税额", "状态"])
            ws.append(["S-1", "2026-01-01", "销项", 900_000, 117_000, "正常"])
            ws.append(["P-1", "2026-01-02", "采购", 500_000, 65_000, "正常"])
            ws = wb.create_sheet("银行流水")
            ws.append(["交易日期", "摘要", "收入金额", "支出金额", "分类"])
            ws.append(["2026-01-02", "销售回款", 100_000, 0, "经营回款"])
            wb.save(path)
            wb.close()
            data = loader.load(path)
            self.assertEqual(str(data.get("利润表.营业收入")), "1280000")
            self.assertEqual(str(data.get("历史.上年同期收入")), "1000000")
            self.assertEqual(str(data.get("发票.销项净额")), "900000")
            self.assertEqual(str(data.get("凭证.本期确认抵扣税额")), "65000")
            self.assertEqual(str(data.get("银行.收入流水合计")), "100000")

    def test_scoring_explains_misses_and_false_positives(self):
        data = loader.load(ROOT / "samples" / "样例企业-审计材料.xlsx")
        findings = engine.run(self.rules, data)
        hits = [f.rule.id for f in findings if f.status == "hit"]
        passed = next(f.rule.id for f in findings if f.status == "pass")
        result = training.score_submission(findings, [hits[0], passed], false_positive_penalty=7)
        self.assertGreater(len(result["missed"]), 0)
        self.assertEqual(len(result["false_positives"]), 1)
        self.assertEqual(result["false_positive_deduction"], 7)
        self.assertIn("误报", result["false_positives"][0]["explanation"])

    def test_pure_synthetic_sample_is_allowed_for_training(self):
        data = loader.load(ROOT / "samples" / "合成测试数据-20260921" / "02-收入少申报-合成.xlsx")
        self.assertTrue(app_module._is_synthetic_dataset(data))


class P1WebFlow(unittest.TestCase):
    def test_setup_persist_audit_assign_and_score(self):
        with tempfile.TemporaryDirectory() as directory:
            db_path = Path(directory) / "web.db"
            old_store = app_module.store
            app_module.store = Store(db_path)
            try:
                with TestClient(app_module.app) as client:
                    setup = client.post("/api/setup", json={
                        "username": "rootadmin", "password": "platform-pass-2026",
                        "display_name": "平台管理员", "org_id": "school-a",
                    })
                    self.assertEqual(setup.status_code, 200, setup.text)
                    self.assertEqual(client.post("/api/login", json={"username": "rootadmin", "password": "platform-pass-2026"}).status_code, 200)
                    for username, role in (("teacher1", "teacher"), ("student1", "student")):
                        app_module.store.create_user(username, f"{username}-pass-2026", username, role, "school-a")
                    client.post("/api/logout")

                    self.assertEqual(client.post("/api/login", json={"username": "teacher1", "password": "teacher1-pass-2026"}).status_code, 200)
                    with (ROOT / "samples" / "样例企业-审计材料.xlsx").open("rb") as stream:
                        response = client.post("/api/audit", files={"file": ("sample.xlsx", stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
                    self.assertEqual(response.status_code, 200, response.text)
                    audit_id = response.json()["audit_id"]
                    self.assertEqual(client.get("/api/audits").json()[0]["id"], audit_id)
                    assignment = client.post("/api/assignments", json={
                        "title": "第一课", "audit_id": audit_id,
                        "false_positive_penalty": 5, "published": True,
                    })
                    self.assertEqual(assignment.status_code, 200, assignment.text)
                    assignment_id = assignment.json()["id"]
                    client.post("/api/logout")

                    self.assertEqual(client.post("/api/login", json={"username": "student1", "password": "student1-pass-2026"}).status_code, 200)
                    self.assertEqual(client.get("/api/audits").status_code, 403)
                    exercise = client.get(f"/api/assignments/{assignment_id}")
                    self.assertNotIn("status", exercise.text)
                    result = client.post(f"/api/assignments/{assignment_id}/submit", json={"selected_rule_ids": ["R-001", "R-006"]})
                    self.assertEqual(result.status_code, 200, result.text)
                    self.assertTrue(result.json()["missed"])
                    self.assertTrue(result.json()["false_positives"])

                # A new repository instance proves results survive app/repository restart.
                reopened = Store(db_path)
                self.assertIsNotNone(reopened.get_audit(audit_id))
            finally:
                app_module.store = old_store

    def test_role_boundaries_and_accountant_export_confirmation(self):
        with tempfile.TemporaryDirectory() as directory:
            old_store = app_module.store
            app_module.store = Store(Path(directory) / "roles.db")
            try:
                with TestClient(app_module.app) as client:
                    client.post("/api/setup", json={
                        "username": "rootadmin", "password": "platform-pass-2026",
                        "display_name": "平台管理员", "org_id": "school-a",
                    })
                    self.assertEqual(client.post("/api/login", json={
                        "username": "rootadmin", "password": "platform-pass-2026",
                    }).status_code, 200)
                    created = {}
                    for username, role in (("orgadmin", "org_admin"), ("accountant1", "accountant"), ("accountant2", "accountant")):
                        created[username] = app_module.store.create_user(username, f"{username}-pass-2026", username, role, "school-a")
                    sample = loader.load(ROOT / "samples" / "样例企业-审计材料.xlsx")
                    self.assertEqual(client.post('/api/login', json={'username':'orgadmin','password':'orgadmin-pass-2026'}).status_code, 200)
                    customer = client.post("/api/clients", json={
                        "name": "仿真客户", "taxpayer_id": sample.company.taxpayer_id,
                        "accountant_id": created["accountant1"]["id"],
                    })
                    self.assertEqual(customer.status_code, 200, customer.text)
                    with (ROOT / "samples" / "样例企业-审计材料.xlsx").open("rb") as stream:
                        response = client.post(
                            "/api/audit", data={"client_id": customer.json()["id"]},
                            files={"file": ("sample.xlsx", stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                        )
                    self.assertEqual(response.status_code, 200, response.text)
                    audit_id = response.json()["audit_id"]
                    client.post("/api/logout")

                    self.assertEqual(client.post("/api/login", json={
                        "username": "orgadmin", "password": "orgadmin-pass-2026",
                    }).status_code, 200)
                    denied = client.post("/api/users", json={
                        "username": "forbidden-teacher", "password": "teacher-pass-2026",
                        "display_name": "越权教师", "role": "teacher", "org_id": "school-a",
                    })
                    self.assertEqual(denied.status_code, 403)
                    client.post("/api/logout")

                    self.assertEqual(client.post("/api/login", json={
                        "username": "accountant2", "password": "accountant2-pass-2026",
                    }).status_code, 200)
                    self.assertEqual(client.get(f"/api/audits/{audit_id}").status_code, 404)
                    client.post("/api/logout")

                    self.assertEqual(client.post("/api/login", json={
                        "username": "accountant1", "password": "accountant1-pass-2026",
                    }).status_code, 200)
                    self.assertEqual(client.get(f"/api/audits/{audit_id}").status_code, 200)
                    self.assertEqual(client.get(f"/api/report/{audit_id}").status_code, 409)
                    report = client.get(f"/api/report/{audit_id}?confirm=true")
                    self.assertEqual(report.status_code, 200, report.text)
                    self.assertEqual(report.headers["content-type"], "application/pdf")
            finally:
                app_module.store = old_store


if __name__ == "__main__":
    unittest.main()
