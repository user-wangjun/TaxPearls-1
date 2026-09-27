from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi.testclient import TestClient
from openpyxl import Workbook

from src import config, engine, materials
from webapp import app as app_module
from webapp.storage import Store
from webapp.access import AccessDenied

ROOT = Path(__file__).resolve().parent.parent
FIXTURES = Path(__file__).parent / "fixtures"
COMPANY = {"name": "仿真导入企业", "taxpayer_id": "TEST-UPLOAD", "industry": "服务业", "period": "2026-01"}
KEYS = {key for rule in engine.load_rules(ROOT / "rules") for key in rule.inputs}


def workbook(sheet, rows, company=None):
    wb = Workbook()
    wb.remove(wb.active)
    if company:
        ws = wb.create_sheet("企业信息")
        ws.append(["项目", "内容"])
        for label, key in zip(config.COMPANY_FIELDS, materials.COMPANY_KEYS):
            ws.append([label, company[key]])
    ws = wb.create_sheet(sheet)
    for row in rows:
        ws.append(row)
    stream = BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def accounts(company=COMPANY, amount=100000):
    return workbook("科目余额表", [config.COL_ACCOUNTS,
        ["6001", "主营业务收入", 0, 0, amount, 0], ["6051", "其他业务收入", 0, 0, 0, 0]], company)


def zip_bytes(files):
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        for name, data in files:
            archive.writestr(name, data)
    return out.getvalue()


def selections(docs, reviewed=True):
    return {d["id"]: {"id": d["id"], "company": d["company"], "rows": deepcopy(d["rows"]), "reviewed": reviewed} for d in docs}


class MaterialParsing(unittest.TestCase):
    def test_pdf_text_tables_zero_page_and_ambiguous_columns(self):
        doc = materials.preview([("税表.pdf", (FIXTURES / "materials-text.pdf").read_bytes())], KEYS)[0]
        self.assertEqual(doc["error"], "")
        self.assertEqual(doc["company"], COMPANY)
        self.assertEqual(doc["page_count"], 2)
        rows = {r["name"]: r for r in doc["rows"]}
        self.assertEqual(rows["增值税.销售额"]["value"], "100000.00")
        self.assertEqual(rows["增值税.进项税额"]["value"], "0")
        self.assertEqual(rows["利润表.营业收入"]["page"], 2)
        self.assertNotIn("200000", [r["value"] for r in doc["rows"]])
        with self.assertRaisesRegex(materials.InputError, "核对"):
            materials.build_dataset([doc], selections([doc], False), {}, KEYS)
        selected = selections([doc])
        selected["0"]["rows"][0]["value"] = "90000"
        data = materials.build_dataset([doc], selected, {}, KEYS)
        self.assertEqual(data.get("增值税.销售额"), 90000)
        self.assertIn("100000.00", data.detail_of("增值税.销售额"))
        self.assertIn("PDF 第 1 页", data.source_of("增值税.销售额"))

    def test_scanned_pdf_manual_entry_and_blank_not_zero(self):
        doc = materials.preview([("扫描.pdf", (FIXTURES / "materials-scanned.pdf").read_bytes())], KEYS)[0]
        self.assertFalse(doc["error"])
        self.assertFalse(doc["rows"])
        self.assertTrue(any("AI 视觉" in w for w in doc["warnings"]))
        selected = selections([doc])
        with self.assertRaisesRegex(materials.InputError, "没有可执行"):
            materials.build_dataset([doc], selected, COMPANY, KEYS)
        selected["0"]["rows"] = [{"name": "增值税.销售额", "value": "0", "page": 1, "detail": "扫描页第 1 栏，元"},
                                   {"name": "增值税.销项税额", "value": "", "page": 1, "detail": "未确认"}]
        data = materials.build_dataset([doc], selected, COMPANY, KEYS)
        self.assertEqual(data.get("增值税.销售额"), 0)
        self.assertIsNone(data.get("增值税.销项税额"))
        self.assertIn("无，人工补录", data.detail_of("增值税.销售额"))
        selected["0"]["rows"][0]["page"] = 2
        with self.assertRaisesRegex(materials.InputError, "页码"):
            materials.build_dataset([doc], selected, COMPANY, KEYS)

    def test_merge_partial_excel_pdf_and_provenance(self):
        docs = materials.preview([("账.xlsx", accounts()), ("税.pdf", (FIXTURES / "materials-text.pdf").read_bytes())], KEYS)
        data = materials.build_dataset(docs, selections(docs), {}, KEYS)
        self.assertEqual(data.get("营业收入"), 100000)
        self.assertEqual(data.get("增值税.销售额"), 100000)
        self.assertIn("账.xlsx", data.source_of("营业收入"))
        self.assertIn("税.pdf", data.source_of("增值税.销售额"))
        findings = engine.run(engine.load_rules(ROOT / "rules"), data)
        self.assertEqual(next(f for f in findings if f.rule.id == "R-001").status, "pass")
        self.assertTrue(any(f.status == "skipped" for f in findings))

    def test_conflicting_company_period_and_metrics_rejected(self):
        for override in ({"name": "另一企业"}, {"taxpayer_id": "OTHER"}, {"period": "2026-02"}):
            docs = materials.preview([("a.xlsx", accounts()), ("b.xlsx", accounts({**COMPANY, **override}))], KEYS)
            with self.assertRaisesRegex(materials.InputError, "不一致"):
                materials.build_dataset(docs, selections(docs), {}, KEYS)
        docs = materials.preview([("a.xlsx", accounts()), ("b.xlsx", accounts(amount=90000))], KEYS)
        with self.assertRaisesRegex(materials.InputError, "冲突"):
            materials.build_dataset(docs, selections(docs), {}, KEYS)
        docs = materials.preview([("a.xlsx", accounts()), ("copy.xlsx", accounts())], KEYS)
        self.assertEqual(materials.build_dataset(docs, selections(docs), {}, KEYS).get("营业收入"), 100000)

    def test_excel_partial_metadata_and_original_validation(self):
        docs = materials.preview([("利润.xlsx", workbook("利润表", [["项目", "本期金额"], ["营业收入", 500]]))], KEYS)
        with self.assertRaisesRegex(materials.InputError, "企业信息"):
            materials.build_dataset(docs, selections(docs), {}, KEYS)
        self.assertEqual(materials.build_dataset(docs, selections(docs), COMPANY, KEYS).get("利润表.营业收入"), 500)
        bad = workbook("利润表", [["项目", "本期金额"], ["营业收入", "=1+1"]], COMPANY)
        self.assertIn("有效数值", materials.preview([("公式.xlsx", bad)], KEYS)[0]["error"])

    def test_zip_bounds_paths_bad_files_and_no_nested_zip(self):
        files = [("目录/a.xlsx", accounts()), ("说明.txt", b"unsupported")]
        docs = materials.preview([("pack.zip", zip_bytes(files))], KEYS)
        self.assertEqual(len(docs), 2)
        self.assertFalse(docs[0]["error"])
        self.assertIn("不支持", docs[1]["error"])
        for files in ([('../a.xlsx', b'x')], [('nested.zip', b'x')], [("bomb.pdf", b'0' * 500000)], [(str(i)+'.pdf', b'x') for i in range(21)]):
            with self.assertRaises(materials.InputError):
                materials.expand_uploads([("bad.zip", zip_bytes(files))])
        for files in ([("a.pdf", b"")], [("a.pdf", b'x'*(materials.MAX_FILE+1))]):
            with self.assertRaises(materials.InputError):
                materials.expand_uploads(files)
        self.assertTrue(materials.preview([("bad.pdf", b'bad')], KEYS)[0]["error"])


class MaterialWebFlow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_store = app_module.store
        app_module.store = Store(Path(self.temp.name) / "test.db")
        app_module.store.create_user("admin", "material-test-2026", "测试管理", "org_admin", "default")
        self.client = TestClient(app_module.app)
        self.client.post("/api/login", json={"username":"admin", "password":"material-test-2026"})

    def tearDown(self):
        self.client.close()
        app_module.store = self.old_store
        self.temp.cleanup()

    def preview(self, files):
        response = self.client.post("/api/materials/preview", files=[("files", (name, data)) for name, data in files])
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def commit(self, draft, **kw):
        return self.client.post("/api/materials/audit", json={"token":draft["token"], "mode":"separate", "selections":list(selections(draft["documents"]).values()), **kw})

    def test_batch_independent_errors_and_idempotency(self):
        draft = self.preview([("a.xlsx", accounts()), ("bad.pdf", b"bad")])
        response = self.commit(draft)
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(len(body["results"]), 1)
        self.assertEqual(len(body["errors"]), 1)
        self.assertEqual(self.commit(draft).json(), body)
        self.assertEqual(len(self.client.get("/api/audits").json()), 1)

    def test_batch_permission_failure_does_not_replay_saved_groups(self):
        draft = self.preview([('a.xlsx',accounts()),('b.xlsx',accounts({**COMPANY,'taxpayer_id':'TEST-B'}))])
        original = app_module.store.save_audit
        count = 0
        def save(*args, **kwargs):
            nonlocal count
            count += 1
            if count==2:
                raise AccessDenied('客户权限已变化。',403)
            return original(*args,**kwargs)
        with patch.object(app_module.store,'save_audit',side_effect=save):
            response = self.commit(draft)
            self.assertEqual(response.status_code,200,response.text)
            result = response.json()
            self.assertEqual(len(result['results']),1)
            self.assertEqual(len(result['errors']),1)
            self.assertEqual(self.commit(draft).json(),result)
        self.assertEqual(count,2)
        self.assertEqual(len(self.client.get('/api/audits').json()),1)

    def test_batch_unknown_after_commit_freezes_result_and_stops_remaining(self):
        draft = self.preview([(name+'.xlsx',accounts({**COMPANY,'taxpayer_id':'TEST-'+name})) for name in ['A','B','C']])
        original = app_module.store.save_audit
        count = 0
        def save(*args, **kwargs):
            nonlocal count
            count += 1
            original(*args,**kwargs)
            if count==2:
                raise RuntimeError('private database path and secret must not escape')
        with patch.object(app_module.store,'save_audit',side_effect=save):
            response = self.commit(draft)
            self.assertEqual(response.status_code,200,response.text)
            result = response.json()
            self.assertEqual(len(result['results']),1)
            self.assertEqual(len(result['errors']),2)
            self.assertIn('无法确认',result['errors'][0]['detail'])
            self.assertIn('尚未执行',result['errors'][1]['detail'])
            self.assertNotIn('private',response.text)
            self.assertEqual(self.commit(draft).json(),result)
        self.assertEqual(count,2)
        self.assertEqual(len(self.client.get('/api/audits').json()),2)

    def test_first_group_unknown_is_not_retried_but_validation_can_be_corrected(self):
        draft = self.preview([('a.xlsx',accounts())])
        with patch.object(app_module.store,'save_audit',side_effect=RuntimeError('local failure')) as save:
            first = self.commit(draft).json()
            self.assertEqual(first['results'],[])
            self.assertEqual(self.commit(draft).json(),first)
            self.assertEqual(save.call_count,1)
        other = self.preview([('no-company.xlsx',accounts(company=None))])
        failed = self.commit(other,mode='merge',same_scope=True).json()
        self.assertEqual(failed['results'],[])
        fixed = self.commit(other,mode='merge',same_scope=True,company=COMPANY).json()
        self.assertEqual(len(fixed['results']),1)

    def test_zip_pdf_merge_review_and_persisted_sources(self):
        draft = self.preview([("pack.zip", zip_bytes([("a.xlsx", accounts()), ("tax.pdf", (FIXTURES/"materials-text.pdf").read_bytes())]))])
        bad = self.commit(draft, mode="merge")
        self.assertEqual(bad.status_code, 422)
        response = self.commit(draft, mode="merge", same_scope=True)
        self.assertEqual(response.status_code, 200, response.text)
        body=response.json()
        self.assertFalse(body["errors"], body)
        audit_id = body["results"][0]["audit"]["audit_id"]
        saved=app_module.store.get_audit(audit_id)
        self.assertIn("pack.zip/tax.pdf", saved["dataset"].source_of("增值税.销售额"))

    def test_auth_owner_and_student_role(self):
        draft = self.preview([("a.xlsx", accounts())])
        self.client.post("/api/logout")
        self.assertEqual(self.commit(draft).status_code, 401)
        app_module.store.create_user("other", "material-test-2026", "其他管理员", "org_admin", "default")
        self.client.post("/api/login", json={"username":"other", "password":"material-test-2026"})
        self.assertEqual(self.commit(draft).status_code, 422)
        self.client.post("/api/logout")
        app_module.store.create_user("student", "material-test-2026", "学生", "student", "default")
        self.client.post("/api/login", json={"username":"student", "password":"material-test-2026"})
        self.assertEqual(self.client.post("/api/materials/preview",files={"files":("a.xlsx", accounts())}).status_code,403)

    def test_upload_limit_is_enforced_before_parser_spills(self):
        response = self.client.post("/api/materials/preview",files={"files":("large.pdf", b'x'*(materials.MAX_FILE+1))})
        self.assertEqual(response.status_code,422)
        self.assertIn("10MB", response.text)


if __name__ == "__main__":
    unittest.main()
