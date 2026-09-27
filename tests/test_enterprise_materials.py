"""Persistent enterprise API integration and confirmation atomicity."""
import base64
from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict, replace
from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from starlette.datastructures import UploadFile

from tests.test_materials import accounts, workbook, zip_bytes, COMPANY
from tests.test_related_graph import graph_workbook
from src import related_graph, config
from src.models import Finding
from webapp import app as module, material_batches
from webapp.storage import Store


class EnterpriseMaterialTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        env = patch.dict(os.environ, {'TAXPEARLS_MATERIAL_KEY': base64.b64encode(b'e' * 32).decode(),
                                     'TAXPEARLS_AI_ENABLED': '0', 'TAXPEARLS_NOTIFICATION_EMAIL_ENABLED': '0'})
        env.start()
        self.addCleanup(env.stop)
        self.store = Store(Path(self.tmp.name) / 'enterprise.db')
        replacement = patch.object(module, 'store', self.store)
        replacement.start()
        self.addCleanup(replacement.stop)
        self.admin = self.store.create_user('ent-admin', 'Enterprise-test-2026!', '企业管理员', 'org_admin', 'ent')
        self.accountant = self.store.create_user('ent-accountant', 'Enterprise-test-2026!', '会计', 'accountant', 'ent')
        self.teacher = self.store.create_user('ent-teacher', 'Enterprise-test-2026!', '教师', 'teacher', 'ent')
        self.foreign = self.store.create_user('ent-other', 'Enterprise-test-2026!', '异机构', 'org_admin', 'other')
        self.client = TestClient(module.app, raise_server_exceptions=False)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.login(self.admin)

    def login(self, actor):
        self.client.cookies.clear()
        token = self.store.authenticate(actor['username'], 'Enterprise-test-2026!')[1]
        self.client.cookies.set(module.COOKIE_NAME, token)

    def upload(self, raw=None):
        response = self.client.post('/api/enterprise/materials', files={'files': ('账.xlsx', raw or accounts())})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def url(self, item, suffix=''):
        return '/api/enterprise/materials/' + item['id'] + suffix

    def confirm(self, item, revision=None):
        return self.client.post(self.url(item, '/confirm'), json={'expected_revision': revision or item['revision']})

    def counts(self):
        with self.store.connect() as db:
            return {table: db.execute('SELECT COUNT(*) FROM ' + table).fetchone()[0] for table in
                    ('audits', 'clients', 'audit_report_versions', 'material_batches', 'material_executions', 'material_revisions')}

    def correct(self, item, changes):
        return self.client.post(self.url(item, '/analyze'), json={'expected_revision': item['revision'],
            'selections': {item['documents'][0]['id']: {'standard_edits': changes}}})

    def operation(self, response):
        response_id = response.headers['x-material-operation']
        rows = self.client.get('/api/enterprise/material-operations')
        self.assertEqual(rows.status_code, 200, rows.text)
        return next(row for row in rows.json() if row['id'] == response_id)

    def test_request_journal_records_pre_batch_failure_without_payload_or_credentials(self):
        with patch('webapp.enterprise_materials.materials.preview') as parse:
            response = self.client.post('/api/enterprise/materials', data={'extraction': 'secret-invalid-mode'},
                                        files={'files': ('private-secret.pdf', b'confidential bytes')})
        self.assertEqual(response.status_code, 422)
        parse.assert_not_called()
        record = self.operation(response)
        self.assertEqual(record['state'], 'failed')
        self.assertIsNone(record['batch_id'])
        self.assertEqual(record['failure_code'], 'invalid_input')
        self.assertEqual(self.counts()['material_batches'], 0)
        with self.store.connect() as db:
            rows = str([tuple(row) for row in db.execute('SELECT * FROM material_operations')])
            logs = str([tuple(row) for row in db.execute("SELECT * FROM audit_log WHERE action LIKE 'material_request_%'")])
        for secret in ('private-secret', 'confidential bytes', 'secret-invalid-mode'):
            self.assertNotIn(secret, rows + logs)

    def test_failed_business_transaction_keeps_failure_journal_and_retry_is_separate(self):
        item = self.upload()
        before = self.counts()
        with self.store.connect() as db:
            db.execute("CREATE TRIGGER reject_confirm_request BEFORE INSERT ON audit_log WHEN NEW.action='material_execution' BEGIN SELECT RAISE(ABORT,'secret'); END")
        failed = self.confirm(item)
        self.assertEqual(failed.status_code, 500)
        self.assertNotIn('secret', failed.text)
        record = self.operation(failed)
        self.assertEqual(record['state'], 'failed')
        self.assertEqual(self.counts(), before)
        with self.store.connect() as db:
            db.execute('DROP TRIGGER reject_confirm_request')
        good = self.confirm(item)
        self.assertEqual(good.status_code, 200, good.text)
        succeeded = self.operation(good)
        self.assertEqual(succeeded['state'], 'committed')
        self.assertNotEqual(succeeded['id'], record['id'])
        self.assertEqual(self.operation(failed), record)

    def test_response_failure_after_commit_never_claims_rollback(self):
        with patch('webapp.enterprise_materials.batches.read', side_effect=RuntimeError('private-exception')):
            response = self.client.post('/api/enterprise/materials', files={'files': ('ledger.xlsx', accounts())})
        self.assertEqual(response.status_code, 500)
        self.assertNotIn('private-exception', response.text)
        record = self.operation(response)
        self.assertEqual(record['state'], 'committed')
        self.assertEqual(record['http_status'], 500)
        self.assertTrue(record['batch_id'])
        self.assertEqual(self.counts()['material_batches'], 1)
        self.assertEqual(self.client.get('/api/enterprise/materials/' + record['batch_id']).status_code, 200)

    def test_journal_start_failure_stops_business_and_observation_failure_stays_unknown(self):
        with self.store.connect() as db:
            db.execute("CREATE TRIGGER fail_request_start BEFORE INSERT ON audit_log WHEN NEW.action='material_request_started' BEGIN SELECT RAISE(ABORT,'secret'); END")
        with patch('webapp.enterprise_materials.materials.preview') as parse:
            response = self.client.post('/api/enterprise/materials', files={'files': ('ledger.xlsx', accounts())})
        self.assertEqual(response.status_code, 503)
        parse.assert_not_called()
        self.assertEqual(self.counts()['material_batches'], 0)
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM material_operations').fetchone()[0], 0)
            db.execute('DROP TRIGGER fail_request_start')
            db.execute("CREATE TRIGGER fail_request_observe BEFORE INSERT ON audit_log WHEN NEW.action='material_request_observed' BEGIN SELECT RAISE(ABORT,'secret'); END")
        response = self.client.post('/api/enterprise/materials', data={'extraction': 'invalid'})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.headers['x-material-operation-state'], 'unknown')
        record = self.operation(response)
        self.assertEqual(record['state'], 'started')
        self.assertIsNone(record['observed_at'])
        with patch.object(module, 'store', Store(self.store.path)):
            self.assertEqual(self.operation(response), record)

    def test_journal_scope_uses_current_client_assignment_and_role(self):
        item = self.upload()
        self.confirm(item)
        with self.store.connect() as db:
            client_id = db.execute('SELECT client_id FROM material_batches WHERE id=?', (item['id'],)).fetchone()[0]
            db.execute('UPDATE clients SET accountant_id=? WHERE id=?', (self.accountant['id'], client_id))
        self.login(self.accountant)
        rows = self.client.get('/api/enterprise/material-operations').json()
        self.assertTrue(any(row['batch_id'] == item['id'] for row in rows))
        with self.store.connect() as db:
            db.execute('UPDATE clients SET accountant_id=NULL WHERE id=?', (client_id,))
        self.assertFalse(self.client.get('/api/enterprise/material-operations').json())
        self.login(self.foreign)
        self.assertFalse(self.client.get('/api/enterprise/material-operations').json())
        self.login(self.teacher)
        self.assertEqual(self.client.get('/api/enterprise/material-operations').status_code, 403)

    def test_repeated_confirmation_has_distinct_committed_requests_one_result(self):
        item = self.upload()
        first, second = self.confirm(item), self.confirm(item)
        self.assertEqual(first.json()['audit_id'], second.json()['audit_id'])
        one, two = self.operation(first), self.operation(second)
        self.assertNotEqual(one['id'], two['id'])
        self.assertEqual([one['state'], two['state']], ['committed', 'committed'])
        self.assertEqual(self.counts()['audits'], 1)

    def test_interrupted_request_is_not_invented_as_failed_after_restart(self):
        from webapp import material_operations
        entry = material_operations.begin(self.store, self.admin, 'upload')
        reopened = Store(self.store.path)
        record = next(r for r in material_operations.listing(reopened, self.admin) if r['id'] == entry['id'])
        self.assertEqual(record['state'], 'started')
        self.assertIsNone(record['http_status'])
        self.assertIsNone(record['observed_at'])

    def test_operation_journal_survives_encrypted_backup_with_original_session(self):
        from webapp import material_operations
        from scripts.ops_db import create_encrypted_backup, restore_encrypted_backup
        self.upload()
        self.client.post('/api/enterprise/materials', data={'extraction': 'invalid'})
        material_operations.begin(self.store, self.admin, 'upload')
        before = self.client.get('/api/enterprise/material-operations').json()
        root = Path(self.tmp.name)
        with patch.dict(os.environ, {'TAXPEARLS_BACKUP_KEY': base64.b64encode(b'b' * 32).decode()}):
            create_encrypted_backup(self.store.path, root / 'journal.tpbackup', retention_days=30)
            restore_encrypted_backup(root / 'journal.tpbackup', root / 'restored.db', safety_retention_days=30)
        with patch.object(module, 'store', Store(root / 'restored.db')):
            response = self.client.get('/api/enterprise/material-operations')
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json(), before)

    def test_missing_commit_journal_rolls_back_business_and_download_failure_is_recorded(self):
        from webapp import material_operations
        real_commit = material_operations.committed
        def lose_journal(db, actor, batch, event):
            if event == 'upload':
                db.execute('DELETE FROM material_operations WHERE id=?', (material_operations.ACTIVE.get()['id'],))
            return real_commit(db, actor, batch, event)
        with patch('webapp.material_operations.committed', side_effect=lose_journal):
            response = self.client.post('/api/enterprise/materials', files={'files': ('ledger.xlsx', accounts())})
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.operation(response)['state'], 'failed')
        self.assertEqual(self.counts()['material_batches'], 0)
        item = self.upload()
        fid = item['files'][0]['id']
        deleted = self.client.request('DELETE', self.url(item, '/originals/' + fid), json={'expected_revision': item['revision']})
        self.assertEqual(deleted.status_code, 200, deleted.text)
        unavailable = self.client.get(self.url(item, '/originals/' + fid))
        self.assertEqual(unavailable.status_code, 410)
        self.assertEqual(self.operation(unavailable)['state'], 'failed')
        self.assertEqual(self.operation(unavailable)['failure_code'], 'deleted')

    def test_revocation_during_parsing_records_failure_without_business_commit(self):
        from src import materials
        real_preview = materials.preview
        def revoke(*args, **kwargs):
            with self.store.connect() as db:
                db.execute('UPDATE users SET active=0 WHERE id=?', (self.admin['id'],))
            return real_preview(*args, **kwargs)
        with patch('webapp.enterprise_materials.materials.preview', side_effect=revoke):
            response = self.client.post('/api/enterprise/materials', files={'files': ('ledger.xlsx', accounts())})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.get('/api/enterprise/material-operations').status_code, 401)
        with self.store.connect() as db:
            row = db.execute('SELECT * FROM material_operations WHERE id=?', (response.headers['x-material-operation'],)).fetchone()
            self.assertEqual(row['state'], 'failed')
            self.assertEqual(row['failure_code'], 'permission')
        self.assertEqual(self.counts()['material_batches'], 0)

    def test_result_material_snapshot_and_report_reference_are_frozen_after_supplement(self):
        item = self.upload()
        original = self.confirm(item).json()
        aid = original['audit_id']
        reference = {'batch_id': item['id'], 'analysis_revision': 1, 'confirmation_revision': 2}
        self.assertEqual(original['material_reference'], reference)
        url = '/api/audits/' + aid + '/materials'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200, response.text)
        snapshot = response.json()['confirmation']
        self.assertEqual(snapshot['audit_id'], aid)
        self.assertNotIn('dataset', snapshot)
        self.assertNotIn('accounts', snapshot['documents'][0])
        archived = self.store.get_report_version(aid, 1)
        self.assertEqual(archived['manifest']['material_reference'], reference)
        self.assertIn('确认材料版本', archived['html'])
        self.assertIn(item['id'], archived['html'])
        item = self.client.get(self.url(item)).json()
        item = self.correct(item, {'科目余额表!E2': '80000'}).json()
        added = self.client.post(self.url(item, '/supplement'), data={'expected_revision': str(item['revision'])},
            files={'files': ('申报.xlsx', workbook('增值税申报', [['项目', '金额'], ['销售额', 80000]], COMPANY))})
        self.assertEqual(added.status_code, 200, added.text)
        newer = self.confirm(added.json()).json()
        self.assertNotEqual(newer['audit_id'], aid)
        module.store = Store(self.store.path)
        old_again = self.client.get(url).json()['confirmation']
        self.assertGreater(old_again.pop('current_revision'), snapshot.pop('current_revision'))
        self.assertEqual(old_again, snapshot)
        self.assertEqual(self.store.get_report_version(aid, 1)['html'], archived['html'])
        self.assertEqual(self.client.get('/api/audits/' + aid).json()['material_reference'], reference)
        self.assertEqual(len(old_again['files']), 1)
        self.assertEqual(len(self.client.get('/api/audits/' + newer['audit_id'] + '/materials').json()['confirmation']['files']), 2)
        trace = self.client.get(self.url(item, '/trace')).json()
        self.assertTrue(any(e['action'] == 'view_confirmation' for e in trace['events']))

    def test_result_material_snapshot_deletion_and_current_permissions(self):
        item = self.upload()
        aid = self.confirm(item).json()['audit_id']
        url = '/api/audits/' + aid + '/materials'
        fid = item['files'][0]['id']
        before = self.client.get(url).json()['confirmation']
        removed = self.client.request(
            'DELETE', self.url(item, '/originals/' + fid), json={'expected_revision': 2})
        self.assertEqual(removed.status_code, 200, removed.text)
        after = self.client.get(url).json()['confirmation']
        self.assertTrue(after['files'][0]['deleted_at'])
        self.assertEqual(after['metrics'], before['metrics'])
        self.assertEqual(after['analysis_sha256'], before['analysis_sha256'])
        self.assertEqual(self.client.get(self.url(item, '/originals/' + fid)).status_code, 410)
        for actor in (self.foreign, self.teacher, self.accountant):
            self.login(actor)
            self.assertIn(self.client.get(url).status_code, (403, 404))
        self.store.upsert_client(self.admin, COMPANY['name'], COMPANY['taxpayer_id'], self.accountant['id'])
        self.assertEqual(self.client.get(url).status_code, 200)
        with self.store.connect() as db:
            db.execute('UPDATE clients SET accountant_id=NULL WHERE org_id=? AND taxpayer_id=?',
                       (self.admin['org_id'], COMPANY['taxpayer_id']))
        self.assertIn(self.client.get(url).status_code, (403, 404))
        self.login(self.admin)
        with self.store.connect() as db:
            db.execute('UPDATE users SET active=0 WHERE id=?', (self.admin['id'],))
        self.assertIn(self.client.get(url).status_code, (401, 403))

    def test_confirmation_read_checks_snapshot_integrity_and_logs_atomically(self):
        item = self.upload()
        aid = self.confirm(item).json()['audit_id']
        url = '/api/audits/' + aid + '/materials'
        with self.store.connect() as db:
            before = db.execute('SELECT COUNT(*) FROM material_events').fetchone()[0]
            db.execute("CREATE TRIGGER fail_confirmation_view BEFORE INSERT ON audit_log WHEN NEW.action='material_view_confirmation' BEGIN SELECT RAISE(ABORT,'forced'); END")
        self.assertEqual(self.client.get(url).status_code, 500)
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM material_events').fetchone()[0], before)
            db.execute('DROP TRIGGER fail_confirmation_view')
            cipher = db.execute('SELECT payload_cipher FROM material_revisions WHERE batch_id=? AND revision=1', (item['id'],)).fetchone()[0]
            db.execute('UPDATE material_revisions SET payload_cipher=? WHERE batch_id=? AND revision=1',
                       (cipher[:-1] + bytes([cipher[-1] ^ 1]), item['id']))
        response = self.client.get(url)
        self.assertEqual(response.status_code, 409, response.text)
        self.assertNotIn('TEST-UPLOAD', response.text)

    def test_report_rejects_wrong_confirmation_reference_and_legacy_does_not_invent_one(self):
        from webapp.report_archive import build_snapshot
        item = self.upload()
        aid = self.confirm(item).json()['audit_id']
        entry = self.store.get_audit(aid)
        changed = deepcopy(entry)
        changed['material_reference']['analysis_revision'] = 99
        wrong = build_snapshot(changed, module._org_branding(self.admin['org_id']))
        with self.assertRaisesRegex(ValueError, '确认版本关联'):
            self.store.archive_report(aid, self.admin, wrong)
        self.assertEqual(len(self.store.report_versions(aid)), 1)
        # A pre-retention result has no link. It must not inherit the most
        # recent batch merely because taxpayer/period happen to match.
        legacy = module._save_audit(entry['dataset'], self.admin)
        self.assertIsNone(legacy['material_reference'])
        self.assertIsNone(self.client.get('/api/audits/' + legacy['audit_id'] + '/materials').json()['confirmation'])

    def test_missing_material_key_blocks_snapshot_not_saved_result_or_report(self):
        item = self.upload()
        aid = self.confirm(item).json()['audit_id']
        saved_html = self.store.get_report_version(aid, 1)['html']
        with patch.dict(os.environ, {'TAXPEARLS_MATERIAL_KEY': ''}):
            self.assertEqual(self.client.get('/api/audits/' + aid + '/materials').status_code, 409)
            self.assertEqual(self.client.get('/api/audits/' + aid).status_code, 200)
            response = self.client.get('/api/report/' + aid + '/html?version=1')
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.text, saved_html)

    def test_confirmation_link_cannot_point_to_another_valid_analysis(self):
        item = self.upload()
        aid = self.confirm(item).json()['audit_id']
        item = self.client.get(self.url(item)).json()
        changed = self.correct(item, {'科目余额表!E2': '80000'}).json()
        with self.store.connect() as db:
            db.execute('UPDATE material_executions SET analysis_revision=? WHERE audit_id=?', (changed['revision'], aid))
        response = self.client.get('/api/audits/' + aid + '/materials')
        self.assertEqual(response.status_code, 409, response.text)
        self.assertNotIn('80000', response.text)

    def test_standard_corrections_recalculate_preserve_original_and_each_revision(self):
        raw = accounts()
        item = self.upload(raw)
        document = item['documents'][0]
        self.assertNotIn('standard_tables', document)
        self.assertNotIn('accounts', document)
        field = next(f for f in document['standard_fields'] if f['id'] == '科目余额表!E2')
        self.assertEqual(field['value'], '100000')
        for amount, old in [('80000', '100000'), ('90000', '80000')]:
            stale = item
            response = self.correct(item, {'科目余额表!E2': amount})
            self.assertEqual(response.status_code, 200, response.text)
            item = response.json()
            metric = next(m for m in item['metrics'] if m['name'] == '营业收入')
            self.assertEqual(str(metric['value']), amount)
            self.assertIn('用户修正', metric['source'])
            self.assertEqual(self.confirm(stale).status_code, 409)
            trace = self.client.get(self.url(item, '/trace')).json()
            version = trace['versions'][-1]
            change = next(c for c in version['detail']['changes'] if c['field'] == '科目余额表!E2')
            self.assertEqual((change['old'], change['new'], change['original']), (old, amount, '100000'))
            self.assertEqual(version['created_by'], self.admin['id'])
            self.assertTrue(version['created_at'])
        first = self.confirm(item)
        self.assertEqual(first.status_code, 200, first.text)
        first_id = first.json()['audit_id']
        frozen = self.client.get('/api/audits/' + first_id).json()
        module.store = Store(self.store.path)
        item = self.client.get(self.url(item)).json()
        self.assertEqual(item['selections'][document['id']]['standard_edits'], {'科目余额表!E2': '90000'})
        restored = self.correct(item, {})
        self.assertEqual(restored.status_code, 200, restored.text)
        result = self.confirm(restored.json())
        self.assertEqual(result.status_code, 200, result.text)
        self.assertNotEqual(result.json()['audit_id'], first_id)
        self.assertEqual(self.client.get('/api/audits/' + first_id).json(), frozen)
        self.assertEqual(self.client.get(self.url(item, '/originals/' + item['files'][0]['id'])).content, raw)

    def test_standard_corrections_survive_supplement_and_partial_purpose_update(self):
        item = self.upload()
        item = self.correct(item, {'科目余额表!E2': '80000'}).json()
        document_id = item['documents'][0]['id']
        added = self.client.post(self.url(item, '/supplement'), data={'expected_revision': str(item['revision'])},
            files={'files': ('申报.xlsx', workbook('增值税申报', [['项目', '金额'], ['销售额', 80000]], COMPANY))})
        self.assertEqual(added.status_code, 200, added.text)
        item = added.json()
        for purpose in ('excluded', 'current'):
            response = self.client.post(self.url(item, '/analyze'), json={'expected_revision': item['revision'],
                'selections': {document_id: {'purpose': purpose}}})
            self.assertEqual(response.status_code, 200, response.text)
            item = response.json()
            self.assertEqual(item['selections'][document_id]['standard_edits'], {'科目余额表!E2': '80000'})
        self.assertEqual(next(m['value'] for m in item['metrics'] if m['name'] == '营业收入'), '80000')
        self.assertEqual(self.confirm(item).status_code, 200)

    def test_standard_invalid_edits_and_log_failure_do_not_commit_partial_state(self):
        item = self.upload()
        for edits in ({'科目余额表!A2': 'bad'}, {'科目余额表!E2': True}, {'科目余额表!E2': '=1+1'},
                      {'科目余额表!E2': {'value': '1', 'source': '伪造'}}):
            self.assertEqual(self.correct(item, edits).status_code, 422)
            self.assertEqual(self.client.get(self.url(item)).json()['revision'], item['revision'])
        before = self.counts()
        with self.store.connect() as db:
            db.execute("CREATE TRIGGER fail_edit BEFORE INSERT ON audit_log WHEN NEW.action='material_edit' BEGIN SELECT RAISE(ABORT,'forced'); END")
        response = self.correct(item, {'科目余额表!E2': '80000'})
        self.assertEqual(response.status_code, 500, response.text)
        self.assertEqual(self.counts(), before)
        self.assertEqual(self.client.get(self.url(item)).json()['selections'], {})
        with self.store.connect() as db:
            db.execute('DROP TRIGGER fail_edit')
        self.assertEqual(self.correct(item, {'科目余额表!E2': '80000'}).status_code, 200)

    def test_standard_domain_error_is_reviewable_and_retry_preserves_missing_not_zero(self):
        item = self.upload(workbook('补充指标', [config.COL_SUPPLEMENT,
            ['人力.社保参保人数', None, '原始底稿', COMPANY['period'], '人数']], COMPANY))
        self.assertFalse(item['analysis']['can_confirm'])
        response = self.correct(item, {'补充指标!B2': '1.5'})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        self.assertFalse(item['analysis']['can_confirm'])
        self.assertEqual(self.confirm(item).status_code, 409)
        self.assertEqual(item['selections'][item['documents'][0]['id']]['standard_edits'], {'补充指标!B2': '1.5'})
        response = self.correct(item, {'补充指标!B2': '0'})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        self.assertTrue(item['analysis']['can_confirm'], item['analysis']['feedback'])
        self.assertEqual(next(m['value'] for m in item['metrics'] if m['name'] == '人力.社保参保人数'), '0')
        item = self.correct(item, {}).json()
        self.assertFalse(item['analysis']['can_confirm'])
        self.assertEqual(item['documents'][0]['standard_fields'][0]['value'], None)

    def test_legacy_zip_sources_upgrade_without_rewriting_old_revision_or_using_ai(self):
        raw = zip_bytes([('账.xlsx', accounts())])
        response = self.client.post('/api/enterprise/materials', files={'files': ('旧材料.zip', raw)})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        view = material_batches.read(self.store, self.admin, item['id'])
        payload = deepcopy(view['versions'][-1]['payload'])
        for doc in payload['documents']:
            for key in ('standard_tables', 'account_cell_sources', 'declaration_cell_sources'):
                doc.pop(key, None)
        revision = material_batches.append_revision(self.store, self.admin, item['id'], item['revision'], 'edit', payload)
        old = self.client.get(self.url(item)).json()
        self.assertIsNone(old['documents'][0]['standard_fields'])
        with patch('src.ai_extraction.AIExtractor.enrich', side_effect=AssertionError('must not call AI')):
            response = self.client.post(self.url(item, '/analyze'), json={'expected_revision': revision})
        self.assertEqual(response.status_code, 200, response.text)
        upgraded = response.json()
        self.assertTrue(upgraded['documents'][0]['standard_fields'])
        view = material_batches.read(self.store, self.admin, item['id'])
        self.assertNotIn('standard_tables', view['versions'][-2]['payload']['documents'][0])
        self.assertTrue(any(e['action'] == 'read_for_reanalysis' for e in view['events']))
        self.assertEqual(self.client.get(self.url(item, '/originals/' + item['files'][0]['id'])).content, raw)

    def test_upload_no_execution_confirm_idempotent_and_restart(self):
        item = self.upload()
        self.assertTrue(item['analysis']['can_confirm'], item)
        self.assertNotIn('dataset', item['analysis'])
        self.assertIn('营业收入', item['fields'])
        self.assertEqual(self.counts()['audits'], 0)
        self.assertEqual(item['revision'], 1)
        result = self.confirm(item)
        self.assertEqual(result.status_code, 200, result.text)
        aid = result.json()['audit_id']
        self.assertEqual(self.confirm(item).json()['audit_id'], aid)
        self.assertEqual(self.counts()['audits'], 1)
        module.store = Store(self.store.path)
        self.assertEqual(self.confirm(item).json()['audit_id'], aid)
        view = self.client.get(self.url(item)).json()
        self.assertEqual(view['revision'], 2)
        self.assertEqual(view['executions'][0]['audit_id'], aid)
        meta = view['files'][0]
        original = self.client.get(self.url(item, '/originals/' + meta['id']))
        self.assertEqual(original.status_code, 200)
        self.assertEqual(original.headers['content-type'], 'application/octet-stream')
        self.assertEqual(original.headers['x-content-type-options'], 'nosniff')

    def test_trace_contains_review_and_confirmation_but_not_raw_datasets(self):
        item = self.upload()
        result = self.confirm(item)
        self.assertEqual(result.status_code, 200, result.text)
        response = self.client.get(self.url(item, '/trace'))
        self.assertEqual(response.status_code, 200, response.text)
        trace = response.json()
        self.assertEqual([v['kind'] for v in trace['versions']], ['analysis', 'confirmation'])
        self.assertNotIn('dataset', trace['versions'][0]['detail']['analysis'])
        self.assertEqual(trace['versions'][1]['detail']['audit_id'], result.json()['audit_id'])
        self.assertEqual(trace['versions'][1]['detail']['scope']['taxpayer_id'], COMPANY['taxpayer_id'])
        self.assertTrue(any(e['action'] == 'execution' for e in trace['events']))
        self.login(self.foreign)
        self.assertIn(self.client.get(self.url(item, '/trace')).status_code, (403, 404))
        self.login(self.teacher)
        self.assertEqual(self.client.get(self.url(item, '/trace')).status_code, 403)

    def test_extraction_failure_trace_is_frozen_encrypted_and_survives_restart_and_retry(self):
        from src.ai_extraction import ExtractionError
        from src.settings import AISettings
        from tests.test_materials import FIXTURES
        import json
        settings = AISettings(enabled=True, api_key='synthetic-provenance-secret', vision=False)
        raw = (FIXTURES / 'materials-text.pdf').read_bytes()
        with patch('webapp.enterprise_materials.AISettings.from_env', return_value=settings), \
                patch('src.ai_extraction.call_model', side_effect=ExtractionError('AI 提取超时', code='timeout')):
            response = self.client.post('/api/enterprise/materials', data={'extraction': 'ai'},
                                        files={'files': ('retry.pdf', raw)})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        trace = self.client.get(self.url(item, '/trace')).json()
        initial = trace['versions'][0]['detail']['extractions'][0]
        self.assertEqual(initial['extraction']['status'], 'failed')
        self.assertEqual(initial['extraction']['failure_code'], 'timeout')
        self.assertEqual(initial['extraction']['calls'], 1)
        self.assertEqual(len(initial['sha256']), 64)
        self.assertNotIn('synthetic-provenance-secret', json.dumps(trace))
        self.assertNotIn('accounts', initial)
        with self.store.connect() as db:
            protected = bytes(db.execute('SELECT payload_cipher FROM material_revisions WHERE batch_id=?',
                                         (item['id'],)).fetchone()[0])
            self.assertNotIn(b'prompt_sha256', protected)
        reopened = Store(self.store.path)
        with patch.object(module, 'store', reopened):
            self.assertEqual(self.client.get(self.url(item, '/trace')).json()['versions'][0]['detail']['extractions'][0], initial)
            retry = self.client.post(self.url(item, '/supplement'), data={'expected_revision': item['revision'], 'extraction': 'local'},
                                     files={'files': ('retry-local.pdf', raw)})
            self.assertEqual(retry.status_code, 200, retry.text)
            updated = self.client.get(self.url(item, '/trace')).json()['versions']
            self.assertEqual(updated[0]['detail']['extractions'][0], initial)
            current = updated[-1]['detail']['extractions']
            self.assertEqual(current[0]['extraction'], initial['extraction'])
            self.assertEqual(current[1]['extraction']['method'], 'local')
            self.assertEqual(current[1]['extraction']['local']['status'], 'succeeded')

    def test_incomplete_company_keeps_failed_file_extraction_trace(self):
        item = self.upload(b'not-a-workbook')
        self.assertFalse(item['analysis']['can_confirm'])
        trace = self.client.get(self.url(item, '/trace')).json()
        row = trace['versions'][0]['detail']['extractions'][0]
        self.assertEqual(row['extraction']['local']['status'], 'failed')
        self.assertTrue(row['error'])
        self.assertEqual(len(row['sha256']), 64)

    def test_edits_invalidate_stale_confirm_and_preserve_prior_revision(self):
        item = self.upload()
        edited = self.client.post(self.url(item, '/analyze'), json={'expected_revision': 1, 'company': {'name': '用户修正名称'}})
        self.assertEqual(edited.status_code, 200, edited.text)
        self.assertEqual(edited.json()['revision'], 2)
        self.assertEqual(self.confirm(item).status_code, 409)
        result = self.confirm(edited.json())
        self.assertEqual(result.status_code, 200, result.text)
        view = material_batches.read(self.store, self.admin, item['id'])
        self.assertEqual(view['versions'][0]['payload']['company']['name'], COMPANY['name'])
        self.assertEqual(view['versions'][1]['payload']['company']['name'], '用户修正名称')

    def test_blocking_identity_or_missing_scope_cannot_confirm(self):
        item = self.upload()
        changed = self.client.post(self.url(item, '/analyze'), json={'expected_revision': 1, 'company': {'taxpayer_id': 'OTHER'}})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertFalse(changed.json()['analysis']['can_confirm'])
        self.assertEqual(self.confirm(changed.json()).status_code, 409)
        self.assertEqual(self.counts()['audits'], 0)

    def test_missing_key_rejects_before_parse_or_model_and_no_rows(self):
        before = self.counts()
        with patch.dict(os.environ, {'TAXPEARLS_MATERIAL_KEY': ''}), \
             patch('src.materials.preview', side_effect=AssertionError('should not parse')):
            result = self.client.post('/api/enterprise/materials', files={'files': ('账.xlsx', accounts())})
            self.assertEqual(result.status_code, 422)
        self.assertEqual(self.counts(), before)

    def test_filling_tax_cannot_hide_unidentified_other_company_then_exclusion_recovers(self):
        item = self.upload(accounts({**COMPANY, 'name': '另一家仿真企业', 'taxpayer_id': ''}))
        changed = self.client.post(self.url(item, '/analyze'), json={'expected_revision': item['revision'],
                          'company': {'name': COMPANY['name'], 'taxpayer_id': COMPANY['taxpayer_id']}})
        self.assertEqual(changed.status_code, 200, changed.text)
        blocked = changed.json()
        self.assertFalse(blocked['analysis']['can_confirm'])
        self.assertIn('identity_unresolved', [n['code'] for n in blocked['analysis']['feedback']['blocking']])
        self.assertEqual(self.confirm(blocked).status_code, 409)
        self.assertEqual(self.counts()['audits'], 0)
        supplemented = self.client.post(self.url(item, '/supplement'),
                        data={'expected_revision': str(blocked['revision'])}, files={'files': ('本企业.xlsx', accounts())})
        self.assertEqual(supplemented.status_code, 200, supplemented.text)
        self.assertFalse(supplemented.json()['analysis']['can_confirm'])
        excluded = self.client.post(self.url(item, '/analyze'), json={
            'expected_revision': supplemented.json()['revision'],
            'selections': {item['documents'][0]['id']: {'purpose': 'excluded'}}})
        self.assertEqual(excluded.status_code, 200, excluded.text)
        self.assertEqual(self.confirm(excluded.json()).status_code, 200)
        stored = material_batches.read(self.store, self.admin, item['id'])
        self.assertEqual(stored['versions'][0]['payload']['documents'][0]['company']['name'], '另一家仿真企业')

    def test_cross_tenant_teacher_and_revocation_denied(self):
        item = self.upload()
        for person in (self.teacher, self.foreign):
            self.login(person)
            self.assertIn(self.client.get(self.url(item)).status_code, (403, 404))
            self.assertIn(self.confirm(item).status_code, (403, 404))
        self.login(self.admin)
        with self.store.connect() as db:
            db.execute('UPDATE users SET active=0 WHERE id=?', (self.admin['id'],))
        self.assertEqual(self.confirm(item).status_code, 401)

    def test_log_failure_rolls_back_confirmation_audit_client_and_report(self):
        item = self.upload()
        before = self.counts()
        with self.store.connect() as db:
            db.execute("CREATE TRIGGER fail_confirm BEFORE INSERT ON audit_log WHEN NEW.action='material_execution' BEGIN SELECT RAISE(ABORT,'forced'); END")
        self.assertEqual(self.confirm(item).status_code, 500)
        self.assertEqual(self.counts(), before)
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT revision FROM material_batches WHERE id=?', (item['id'],)).fetchone()[0], 1)
            db.execute('DROP TRIGGER fail_confirm')
        self.assertEqual(self.confirm(item).status_code, 200)

    def test_concurrent_confirmation_exactly_one_audit(self):
        item = self.upload()
        cookie = self.client.cookies.get(module.COOKIE_NAME)

        def submit(_index):
            with TestClient(module.app, raise_server_exceptions=False) as client:
                client.cookies.set(module.COOKIE_NAME, cookie)
                return client.post(self.url(item, '/confirm'), json={'expected_revision': 1})

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(submit, range(2)))
        self.assertEqual([r.status_code for r in results], [200, 200], [r.text for r in results])
        self.assertEqual(results[0].json()['audit_id'], results[1].json()['audit_id'])
        self.assertEqual(self.counts()['audits'], 1)

    def test_deletion_impact_lists_results_and_blocks_stale_candidate(self):
        item = self.upload()
        result = self.confirm(item)
        self.assertEqual(result.status_code, 200, result.text)
        fid = item['files'][0]['id']
        impact = self.client.get(self.url(item, '/originals/' + fid + '/deletion-impact')).json()
        self.assertEqual(impact['audits'][0]['audit_id'], result.json()['audit_id'])
        deleted = self.client.request('DELETE', self.url(item, '/originals/' + fid), json={'expected_revision': 2})
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(self.client.get(self.url(item, '/originals/' + fid)).status_code, 410)
        self.assertEqual(self.confirm(item, 3).status_code, 409)
        self.assertEqual(self.client.post(self.url(item, '/analyze'), json={'expected_revision': 3}).status_code, 409)

    def test_browser_cannot_supply_dataset_rules_or_confirmation(self):
        item = self.upload()
        for field in ('dataset', 'rules', 'graph_rule', 'analysis', 'confirmation'):
            response = self.client.post(self.url(item, '/confirm'), json={'expected_revision': 1, field: {}})
            self.assertEqual(response.status_code, 422)
        self.assertEqual(self.counts()['audits'], 0)

    def test_confirmation_scope_matches_all_results_including_missing_graph(self):
        item = self.upload()
        graph_check = next(c for c in item['analysis']['checks'] if c['rule_id'] == 'G-001')
        self.assertFalse(graph_check['ready'])
        response = self.confirm(item)
        self.assertEqual(response.status_code, 200, response.text)
        findings = response.json()['findings']
        self.assertEqual({c['rule_id'] for c in item['analysis']['checks']}, {f['id'] for f in findings})
        graph = next(f for f in findings if f['id'] == 'G-001')
        self.assertEqual(graph['status'], 'skipped')
        self.assertEqual(graph['skip_reason'], graph_check['reasons'][0])
        html = self.client.get('/api/report/' + response.json()['audit_id'] + '/html').text
        self.assertIn('G-001', html)
        self.assertIn(graph_check['reasons'][0], html)
        self.assertIn('本次检查范围', html)
        self.assertNotIn('<td>执行规则数</td>', html)
        self.assertNotIn('项检查因资料不足未能执行', html)
        self.assertNotIn('因被审计单位未提供相应资料而未能执行', html)

    def test_graph_definition_change_requires_reanalysis_but_preserves_old_results(self):
        book = graph_workbook()
        stream = BytesIO()
        book.save(stream)
        book.close()
        item = self.upload(stream.getvalue())
        original_definition = related_graph.definition()
        stored = material_batches.read(self.store, self.admin, item['id'])['versions'][0]['payload']
        self.assertEqual(stored['graph_rule'], asdict(original_definition))
        result = self.confirm(item)
        self.assertEqual(result.status_code, 200, result.text)
        graph = next(f for f in result.json()['findings'] if f['id'] == 'G-001')
        self.assertEqual((graph['version'], graph['status']), (original_definition.version, 'hit'))
        pending = self.upload(stream.getvalue())
        changed = replace(original_definition, version='2.0', description='新的图检查口径')
        with patch.object(related_graph, 'definition', return_value=changed):
            with patch.object(related_graph, 'run', side_effect=AssertionError('must reject before execution')):
                self.assertEqual(self.confirm(pending).status_code, 409)
                self.assertEqual(self.confirm(item).json(), result.json())
            refreshed = self.client.post(self.url(pending, '/analyze'), json={'expected_revision': pending['revision']})
            self.assertEqual(refreshed.status_code, 200, refreshed.text)
            retry = self.confirm(refreshed.json())
            self.assertEqual(retry.status_code, 200, retry.text)
            new_graph = next(f for f in retry.json()['findings'] if f['id'] == 'G-001')
            self.assertEqual(new_graph['version'], '2.0')
        module.store = Store(self.store.path)
        old = self.client.get('/api/audits/' + result.json()['audit_id'])
        self.assertEqual(old.json(), result.json())

    def test_execution_rejects_extra_missing_duplicate_or_changed_graph_rules_atomically(self):
        item = self.upload()
        definition = related_graph.definition()
        finding = Finding(definition, 'skipped', None, '', '', skip_reason='缺少图材料')
        variants = [[], [finding, finding], [finding, replace(finding, rule=replace(definition, id='G-999'))],
                    [replace(finding, rule=replace(definition, version='9.0'))]]
        before = self.counts()
        for findings in variants:
            with self.subTest(rules=[f.rule.id for f in findings]), patch.object(related_graph, 'run', return_value=findings):
                response = self.confirm(item)
                self.assertEqual(response.status_code, 409, response.text)
                self.assertEqual(self.counts(), before)
        self.assertEqual(self.confirm(item).status_code, 200)

    def test_old_analysis_without_graph_snapshot_requires_reanalysis(self):
        item = self.upload()
        stored = material_batches.read(self.store, self.admin, item['id'])['versions'][0]['payload']
        stored.pop('graph_rule')
        # Simulate a retained pre-upgrade analysis through the trusted store,
        # not a public endpoint accepting arbitrary snapshot content.
        revision = material_batches.append_revision(self.store, self.admin, item['id'], item['revision'], 'edit', stored)
        self.assertEqual(self.confirm(item, revision).status_code, 409)
        self.assertEqual(self.counts()['audits'], 0)
        refreshed = self.client.post(self.url(item, '/analyze'), json={'expected_revision': revision})
        self.assertEqual(refreshed.status_code, 200, refreshed.text)
        self.assertEqual(self.confirm(refreshed.json()).status_code, 200)

    def test_legacy_enterprise_entrypoints_denied_before_parsing(self):
        for actor in (self.admin, self.accountant, self.foreign):
            self.login(actor)
            with self.subTest(role=actor['role']), \
                 patch.object(UploadFile, 'write', side_effect=AssertionError('must not read upload')), \
                 patch('src.materials.preview', side_effect=AssertionError('must not parse')), \
                 patch('src.engine.run', side_effect=AssertionError('must not execute')):
                for path in ('/api/audit', '/api/materials/preview'):
                    response = self.client.post(path, files={'file': ('untrusted.xlsx', b'not parsed')})
                    self.assertEqual(response.status_code, 403, response.text)
                self.assertEqual(self.client.post('/api/materials/audit', json={
                    'token': 'old-token', 'mode': 'merge', 'same_scope': True,
                    'material_context': {'revision': 1}, 'role': 'teacher'}).status_code, 403)
                self.assertEqual(self.client.get('/api/materials/jobs/old-job').status_code, 403)
        self.assertEqual(self.counts()['audits'], 0)
        self.assertEqual(self.counts()['material_batches'], 0)

    def test_teaching_import_keeps_separate_synthetic_gate(self):
        self.login(self.teacher)
        def complete_book(company):
            book = load_workbook(BytesIO(accounts(company=company)))
            sheet = book.create_sheet('增值税申报')
            sheet.append(['项目', '金额'])
            sheet.append(['销售额', 100000])
            output = BytesIO()
            book.save(output)
            book.close()
            return output.getvalue()

        good = self.client.post('/api/audit', files={'file': ('teaching.xlsx', complete_book(COMPANY))})
        self.assertEqual(good.status_code, 200, good.text)
        real = self.client.post('/api/audit', files={'file': ('business.xlsx', complete_book({
            **COMPANY, 'name': '某真实客户', 'taxpayer_id': '913100001234567890'}))})
        self.assertEqual(real.status_code, 403, real.text)
        self.assertEqual(self.counts()['audits'], 1)
        self.assertEqual(self.counts()['material_batches'], 0)

    def test_supplement_keeps_user_values_old_originals_and_links_new_audit(self):
        item = self.upload()
        original_audit = self.confirm(item).json()['audit_id']
        edit = self.client.post(self.url(item, '/analyze'), json={'expected_revision': 2,
                                'company': {'name': '用户保留的企业名称'}}).json()
        # Same filename is a distinct retained upload with a stable file ID.
        added = self.client.post(self.url(item, '/supplement'), data={'expected_revision': str(edit['revision'])},
                                 files={'files': ('账.xlsx', accounts(COMPANY, 200000))})
        self.assertEqual(added.status_code, 200, added.text)
        new = added.json()
        self.assertEqual(new['company']['name'], '用户保留的企业名称')
        self.assertTrue(new['analysis']['supplement_differences'])
        self.assertEqual(len(new['files']), 2)
        self.assertEqual(len(set(d['original_id'] for d in new['documents'])), 2)
        self.assertFalse(new['analysis']['can_confirm'])  # conflicting ledgers not added together
        self.assertEqual(self.confirm(new).status_code, 409)
        old_document = item['documents'][0]['id']
        revised = self.client.post(self.url(item, '/analyze'), json={'expected_revision': new['revision'],
                         'selections': {old_document: {'purpose': 'excluded'}}})
        self.assertEqual(revised.status_code, 200, revised.text)
        result = self.confirm(revised.json())
        self.assertEqual(result.status_code, 200, result.text)
        self.assertNotEqual(result.json()['audit_id'], original_audit)
        self.assertEqual(self.counts()['audits'], 2)
        self.assertEqual(len(self.client.get(self.url(item)).json()['executions']), 2)

    def test_supplement_log_failure_and_stale_revision_leave_no_partial_files(self):
        item = self.upload()
        with self.store.connect() as db:
            db.execute("CREATE TRIGGER fail_supplement BEFORE INSERT ON audit_log WHEN NEW.action='material_supplement' BEGIN SELECT RAISE(ABORT,'forced'); END")
        for revision, expected in [('0', 409), ('1', 500)]:
            response = self.client.post(self.url(item, '/supplement'), data={'expected_revision': revision},
                                        files={'files': ('补传.xlsx', accounts())})
            self.assertEqual(response.status_code, expected, response.text)
        view = self.client.get(self.url(item)).json()
        self.assertEqual(len(view['files']), 1)
        self.assertEqual(view['revision'], 1)

    def test_assigned_client_rejected_before_parsing_and_accountant_recovery(self):
        cid = self.store.upsert_client(self.admin, '受控企业', COMPANY['taxpayer_id'])
        self.login(self.accountant)
        with patch('src.materials.preview', side_effect=AssertionError('unauthorized parsing')):
            response = self.client.post('/api/enterprise/materials', data={'client_id': cid['id']},
                                        files={'files': ('账.xlsx', accounts())})
            self.assertEqual(response.status_code, 404, response.text)
        self.store.upsert_client(self.admin, '受控企业', COMPANY['taxpayer_id'], self.accountant['id'])
        item = self.upload()
        self.assertEqual(self.client.get('/api/enterprise/materials').json()[0]['id'], item['id'])
        self.assertEqual(self.confirm(item).status_code, 200)
        self.login(self.foreign)
        self.assertEqual(self.client.get('/api/enterprise/materials').json(), [])


if __name__ == '__main__':
    unittest.main()
