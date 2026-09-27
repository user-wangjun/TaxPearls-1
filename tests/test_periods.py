from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from fastapi.testclient import TestClient
from openpyxl import Workbook

from src import config, engine, loader, materials
from webapp import app as app_module
from webapp.storage import Store


ROOT = Path(__file__).resolve().parent.parent
RULES = engine.load_rules(ROOT / "rules")
RULE_KEYS = {key for rule in RULES for key in rule.inputs}


def period_workbook(
    period: str,
    history_rows: list[list[object]],
    *,
    income: list[tuple[str, object]] | None = None,
    balance: list[tuple[str, object]] | None = None,
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet(config.SHEET_COMPANY)
    ws.append(["项目", "内容"])
    for key, value in (
        ("企业名称", "多期间仿真企业"),
        ("纳税人识别号", "PERIOD-TEST-001"),
        ("所属行业", "软件服务业"),
        ("所属期", period),
    ):
        ws.append([key, value])

    ws = wb.create_sheet(config.SHEET_ACCOUNTS)
    ws.append(config.COL_ACCOUNTS)
    for code, name in (
        ("6001", "主营业务收入"),
        ("6051", "其他业务收入"),
        ("6401", "主营业务成本"),
        ("6402", "其他业务成本"),
        ("22210105", "销项税额"),
        ("22210101", "进项税额"),
    ):
        ws.append([code, name, 0, 0, 0, 0])

    ws = wb.create_sheet(config.SHEET_DECLARATION)
    ws.append(["项目", "金额"])

    if income:
        ws = wb.create_sheet(config.SHEET_INCOME)
        ws.append(config.COL_STATEMENT)
        for row in income:
            ws.append(row)
    if balance:
        ws = wb.create_sheet(config.SHEET_BALANCE)
        ws.append(config.COL_STATEMENT)
        for row in balance:
            ws.append(row)

    ws = wb.create_sheet(config.SHEET_HISTORY)
    ws.append(config.COL_HISTORY)
    for row in history_rows:
        ws.append(row)

    stream = BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def history_row(metric: str, value: object, period: str) -> list[object]:
    return [metric, value, f"{period} 月结报表", period, "同一主体、人民币元、完整自然期间"]


def partial_history_workbook(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = config.SHEET_HISTORY
    ws.append(config.COL_HISTORY)
    for row in rows:
        ws.append(row)
    stream = BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


class PeriodAggregationTests(unittest.TestCase):
    def test_monthly_yoy_mom_r12_and_rule_inputs(self):
        rows: list[list[object]] = []
        for month in range(1, 13):
            period = f"2025-{month:02d}"
            rows.append(history_row("利润表.营业收入", 80 if month == 1 else 100, period))
            rows.append(history_row("利润表.营业成本", 40 if month == 1 else 50, period))
        data = loader.load_bytes(period_workbook(
            "2026-01",
            rows,
            income=[("营业收入", 120), ("营业成本", 60)],
        ))

        self.assertEqual(data.get("趋势.营业收入.本期"), Decimal("120"))
        self.assertEqual(data.get("趋势.营业收入.上期"), Decimal("100"))
        self.assertEqual(data.get("趋势.营业收入.上年同期"), Decimal("80"))
        self.assertEqual(data.get("趋势.营业收入.环比率"), Decimal("0.2"))
        self.assertEqual(data.get("趋势.营业收入.同比率"), Decimal("0.5"))
        self.assertEqual(data.get("趋势.营业收入.滚动12月"), Decimal("1220"))
        self.assertEqual(data.get("趋势.营业成本.滚动12月"), Decimal("610"))
        self.assertEqual(data.get("历史.上年同期收入"), Decimal("80"))
        self.assertEqual(data.get("历史.上年同期成本"), Decimal("40"))
        self.assertIn("2025-02", data.detail_of("趋势.营业收入.滚动12月"))
        self.assertIn("2026-01", data.detail_of("趋势.营业收入.滚动12月"))

        findings = {finding.rule.id: finding for finding in engine.run(RULES, data)}
        self.assertEqual(findings["R-020"].status, "pass")
        self.assertEqual(findings["R-022"].status, "hit")

    def test_gap_blocks_partial_r12_and_zero_base_keeps_real_zero(self):
        rows = [history_row("利润表.营业收入", 10, f"2025-{month:02d}") for month in range(3, 13) if month != 6]
        rows.extend([
            history_row("利润表.营业收入", 0, "2025-02"),
            history_row("利润表.营业收入", 0, "2026-01"),
        ])
        data = loader.load_bytes(period_workbook("2026-02", rows, income=[("营业收入", 100)]))

        self.assertEqual(data.get("趋势.营业收入.上期"), Decimal("0"))
        self.assertEqual(data.get("趋势.营业收入.上年同期"), Decimal("0"))
        self.assertEqual(data.get("历史.上年同期收入"), Decimal("0"))
        self.assertIsNone(data.get("趋势.营业收入.环比率"))
        self.assertIsNone(data.get("趋势.营业收入.同比率"))
        self.assertIsNone(data.get("趋势.营业收入.滚动12月"))
        finding = next(finding for finding in engine.run(RULES, data) if finding.rule.id == "R-022")
        self.assertEqual(finding.status, "skipped")
        self.assertIn("基数", finding.skip_reason)

    def test_duplicate_invalid_period_and_current_conflict_are_rejected(self):
        cases = (
            (
                [history_row("利润表.营业收入", 1, "2025-1"), history_row("利润表.营业收入", 1, "2025-01")],
                "重复所属期",
            ),
            ([history_row("利润表.营业收入", 1, "2025-13")], "月份无效"),
        )
        for rows, message in cases:
            with self.subTest(message=message), self.assertRaisesRegex(loader.InputError, message):
                loader.load_bytes(period_workbook("2026-01", rows, income=[("营业收入", 2)]))

        with self.assertRaisesRegex(loader.InputError, "本期值与本期报表不一致"):
            loader.load_bytes(period_workbook(
                "2026-01",
                [history_row("利润表.营业收入", 999, "2026-01")],
                income=[("营业收入", 100)],
            ))

    def test_full_year_series_generates_annual_rule_inputs_without_summing_assets(self):
        rows = [
            history_row("利润表.净利润", -90, "2025"),
            history_row("利润表.净利润", -80, "2024"),
            history_row("资产负债表.资产总计", 140, "2025"),
            history_row("资产负债表.资产总计", 100, "2024"),
        ]
        data = loader.load_bytes(period_workbook(
            "2026",
            rows,
            income=[("净利润", -100)],
            balance=[("资产总计", 180)],
        ))

        expected = {
            "年度.本年净利润": Decimal("-100"),
            "年度.上年净利润": Decimal("-90"),
            "年度.前年净利润": Decimal("-80"),
            "年度.本年末资产": Decimal("180"),
            "年度.上年末资产": Decimal("140"),
            "年度.前年末资产": Decimal("100"),
        }
        self.assertEqual({key: data.get(key) for key in expected}, expected)
        self.assertIsNone(data.get("趋势.资产总计.滚动12月"))
        finding = next(finding for finding in engine.run(RULES, data) if finding.rule.id == "R-024")
        self.assertEqual(finding.status, "hit")

    def test_web_material_preview_and_commit_keep_derived_period_evidence(self):
        rows = [history_row("利润表.营业收入", 80, "2025-01"), history_row("利润表.营业收入", 100, "2025-12")]
        book = period_workbook("2026-01", rows, income=[("营业收入", 120)])
        with tempfile.TemporaryDirectory() as directory:
            old_store = app_module.store
            app_module.store = Store(Path(directory) / "periods.db")
            try:
                app_module.store.create_user("admin", "period-test-2026", "期间测试管理员", "org_admin", "default")
                with TestClient(app_module.app) as client:
                    self.assertEqual(client.post("/api/login", json={"username": "admin", "password": "period-test-2026"}).status_code, 200)
                    preview = client.post("/api/materials/preview", files={"files": ("periods.xlsx", book)}).json()
                    document = preview["documents"][0]
                    rows_by_name = {row["name"]: row for row in document["rows"]}
                    self.assertEqual(rows_by_name["趋势.营业收入.同比率"]["value"], "0.5")
                    selections = [{"id": document["id"], "company": document["company"], "rows": document["rows"], "reviewed": True}]
                    committed = client.post("/api/materials/audit", json={
                        "token": preview["token"], "mode": "separate", "selections": selections,
                    })
                    self.assertEqual(committed.status_code, 200, committed.text)
                    audit_id = committed.json()["results"][0]["audit"]["audit_id"]
                    audit = client.get(f"/api/audits/{audit_id}").json()
                    metrics = {metric["name"]: metric for metric in audit["metrics"]}
                    self.assertEqual(metrics["历史.上年同期收入"]["value"], "80.00")
                    self.assertIn("实际期间 2025-01", metrics["历史.上年同期收入"]["source"])
            finally:
                app_module.store = old_store

    def test_separate_current_and_history_workbooks_are_merged_before_derivation(self):
        current = period_workbook("2026-01", [], income=[("营业收入", 120)])
        history = partial_history_workbook([
            history_row("利润表.营业收入", 80, "2025-01"),
            history_row("利润表.营业收入", 100, "2025-12"),
        ])
        docs = materials.preview([("current.xlsx", current), ("history.xlsx", history)], RULE_KEYS)
        self.assertFalse([doc["error"] for doc in docs if doc["error"]], docs)
        self.assertEqual(len(docs[1]["period_series"]), 2)
        selections = {
            doc["id"]: {"id": doc["id"], "company": doc["company"], "rows": doc["rows"], "reviewed": True}
            for doc in docs
        }
        data = materials.build_dataset(docs, selections, {}, RULE_KEYS)
        self.assertEqual(data.get("趋势.营业收入.环比率"), Decimal("0.2"))
        self.assertEqual(data.get("趋势.营业收入.同比率"), Decimal("0.5"))
        self.assertIn("history.xlsx", data.source_of("历史.上年同期收入"))


if __name__ == "__main__":
    unittest.main()
