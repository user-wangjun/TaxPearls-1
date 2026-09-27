from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from src import config, engine, materials
from webapp import app as app_module
from webapp.storage import Store


ROOT = Path(__file__).resolve().parent.parent
COMPANY = {
    "name": "四流勾稽测试企业",
    "taxpayer_id": "CONTRACT-TEST-001",
    "industry": "批发业",
    "period": "2026H1",
}
KEYS = {key for rule in engine.load_rules(ROOT / "rules") for key in rule.inputs}


def selections(docs):
    return {
        doc["id"]: {
            "id": doc["id"],
            "company": deepcopy(doc["company"]),
            "rows": deepcopy(doc["rows"]),
            "reviewed": True,
        }
        for doc in docs
    }


def contract_workbook(
    *,
    contract_kind="销售",
    contract_amount=1130,
    invoice_kind="销项",
    invoice_amount=1000,
    invoice_tax=130,
    bank_income=1130,
    bank_expense=None,
    fulfillment_kind="发货",
    fulfillment_amount=1130,
    fulfillment_date="2026-02-10",
    fulfillment_review="已复核",
    link_amount=1130,
    link_review="已复核",
    include_invoice=True,
    include_bank=True,
    include_fulfillment=True,
    include_link=True,
):
    workbook = Workbook()
    ws = workbook.active
    ws.title = config.SHEET_COMPANY
    ws.append(["项目", "内容"])
    for label, key in zip(config.COMPANY_FIELDS, COMPANY):
        ws.append([label, COMPANY[key]])

    if include_invoice:
        ws = workbook.create_sheet(config.SHEET_INVOICES)
        ws.append(config.COL_INVOICES)
        ws.append(["INV-001", "2026-02-05", invoice_kind, invoice_amount, invoice_tax, "正常"])

    if include_bank:
        ws = workbook.create_sheet("银行来源")
        ws.append(["唯一编号", "交易日期", "收入金额", "支出金额", "摘要", "币种"])
        ws.append(["TX-001", "2026-02-08", bank_income, bank_expense, "合同款", "CNY"])

    ws = workbook.create_sheet(config.SHEET_CONTRACTS)
    ws.append(config.COL_CONTRACTS)
    ws.append([
        "CT-001", contract_kind, "客户甲", "2026-01-15", contract_amount,
        "2026-01-01", "2026-06-30", "已复核", "WP-C-001", "合同原件已核对",
    ])

    if include_fulfillment:
        ws = workbook.create_sheet(config.SHEET_FULFILLMENTS)
        ws.append(config.COL_FULFILLMENTS)
        ws.append([
            "FUL-001", "CT-001", fulfillment_date, fulfillment_kind, fulfillment_amount,
            fulfillment_review, "WP-F-001", "签收或验收证据已核对" if fulfillment_review == "已复核" else "",
        ])

    if include_link:
        ws = workbook.create_sheet(config.SHEET_CONTRACT_LINKS)
        ws.append(config.COL_CONTRACT_LINKS)
        ws.append([
            "LINK-001", "CT-001", "INV-001" if include_invoice else None,
            "TX-001" if include_bank else None, "FUL-001" if include_fulfillment else None,
            link_amount, link_review, "WP-L-001", "四流分摊已核对" if link_review == "已复核" else "",
        ])

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


class ContractLedgerParsing(unittest.TestCase):
    def test_complete_four_flow_generates_evidence_metrics(self):
        doc = materials.preview([("四流台账.xlsx", contract_workbook())], KEYS)[0]
        self.assertFalse(doc["error"], doc)
        self.assertEqual(len(doc["contracts"]), 1)
        self.assertEqual(len(doc["fulfillments"]), 1)
        self.assertEqual(len(doc["contract_links"]), 1)
        self.assertIn("1 份合同", doc["summary"])

        data = materials.build_dataset([doc], selections([doc]), {}, KEYS)
        self.assertEqual(data.get("合同.合同数量"), 1)
        self.assertEqual(data.get("合同.合同含税金额"), 1130)
        self.assertEqual(data.get("合同.销售合同含税金额"), 1130)
        self.assertIsNone(data.get("合同.采购合同含税金额"))
        self.assertEqual(data.get("合同.四流完整合同数量"), 1)
        self.assertEqual(data.get("合同.四流完整勾稽金额"), 1130)
        self.assertEqual(data.get("合同.四流待完善合同数量"), 0)
        self.assertIn("CT-001", data.detail_of("合同.四流完整勾稽金额"))

    def test_contract_can_be_registered_before_all_flows_arrive(self):
        doc = materials.preview([
            ("合同登记.xlsx", contract_workbook(
                include_invoice=False, include_bank=False, include_fulfillment=False, include_link=False,
            ))
        ], KEYS)[0]
        self.assertFalse(doc["error"], doc)
        data = materials.build_dataset([doc], selections([doc]), {}, KEYS)
        self.assertEqual(data.get("合同.合同数量"), 1)
        self.assertEqual(data.get("合同.四流完整合同数量"), 0)
        self.assertEqual(data.get("合同.四流待完善合同数量"), 1)

    def test_direction_missing_reference_and_overallocation_are_rejected(self):
        wrong_direction = materials.preview([
            ("方向错误.xlsx", contract_workbook(contract_kind="采购", bank_income=None, bank_expense=1130,
                                                 fulfillment_kind="收货"))
        ], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "发票方向"):
            materials.build_dataset([wrong_direction], selections([wrong_direction]), {}, KEYS)

        missing = materials.preview([("缺发票.xlsx", contract_workbook(include_invoice=False))], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "不存在的发票"):
            tampered = deepcopy(missing)
            tampered["contract_links"][0]["invoice_number"] = "INV-MISSING"
            materials.build_dataset([tampered], selections([tampered]), {}, KEYS)

        over = materials.preview([
            ("超额.xlsx", contract_workbook(contract_amount=1200, bank_income=1200,
                                             fulfillment_amount=1200, link_amount=1200))
        ], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "发票.*超过价税合计"):
            materials.build_dataset([over], selections([over]), {}, KEYS)

    def test_invoice_direction_can_be_inferred_from_taxpayer_ids(self):
        workbook = load_workbook(BytesIO(contract_workbook()))
        del workbook[config.SHEET_INVOICES]
        ws = workbook.create_sheet("电子税务局发票查询")
        ws.append([
            "发票号码", "开票日期", "不含税金额", "税额", "发票状态",
            "销方纳税人识别号", "购方纳税人识别号",
        ])
        ws.append(["INV-001", "2026-02-05", 1000, 130, "正常", COMPANY["taxpayer_id"], "BUYER-001"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        doc = materials.preview([("税号推断.xlsx", output.getvalue())], KEYS)[0]
        self.assertFalse(doc["error"], doc)
        self.assertEqual(doc["invoices"][0]["kind"], "销项")
        data = materials.build_dataset([doc], selections([doc]), {}, KEYS)
        self.assertEqual(data.get("合同.四流完整合同数量"), 1)

    def test_fulfillment_period_kind_and_review_gate_completion(self):
        pending = materials.preview([
            ("待复核.xlsx", contract_workbook(fulfillment_review="待复核", link_review="待复核"))
        ], KEYS)[0]
        data = materials.build_dataset([pending], selections([pending]), {}, KEYS)
        self.assertEqual(data.get("合同.四流完整合同数量"), 0)

        inconsistent = materials.preview([
            ("复核状态冲突.xlsx", contract_workbook(fulfillment_review="待复核"))
        ], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "履约单据.*尚未复核"):
            materials.build_dataset([inconsistent], selections([inconsistent]), {}, KEYS)

        contract_pending = materials.preview([("合同待复核.xlsx", contract_workbook())], KEYS)[0]
        contract_pending["contracts"][0]["reviewed"] = False
        with self.assertRaisesRegex(materials.InputError, "合同.*尚未复核"):
            materials.build_dataset([contract_pending], selections([contract_pending]), {}, KEYS)

        wrong_kind = materials.preview([
            ("履约方向.xlsx", contract_workbook(fulfillment_kind="收货"))
        ], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "不能关联收货"):
            materials.build_dataset([wrong_kind], selections([wrong_kind]), {}, KEYS)

        outside = materials.preview([
            ("履约跨期.xlsx", contract_workbook(fulfillment_date="2025-12-31"))
        ], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "不在核对期间"):
            materials.build_dataset([outside], selections([outside]), {}, KEYS)

    def test_cross_file_merge_and_duplicate_link_protection(self):
        primary = materials.preview([("四流证据.xlsx", contract_workbook())], KEYS)[0]
        evidence = deepcopy(primary)
        ledger = deepcopy(primary)
        evidence.update(id="0", name="发票与资金.xlsx", fingerprint="evidence-only")
        ledger.update(id="1", name="合同与履约.xlsx", fingerprint="ledger-only")
        for field in ("contracts", "fulfillments", "contract_links"):
            evidence[field] = []
        for field in ("invoices", "bank_transactions", "bank_adjustments", "human_records"):
            ledger[field] = []
        merged = materials.build_dataset([evidence, ledger], selections([evidence, ledger]), COMPANY, KEYS)
        self.assertEqual(merged.get("合同.四流完整合同数量"), 1)
        source = merged.source_of("合同.四流完整勾稽金额")
        self.assertIn("发票与资金.xlsx", source)
        self.assertIn("合同与履约.xlsx", source)

        duplicate = deepcopy(primary)
        duplicate.update(id="1", name="重复勾稽.xlsx", fingerprint="duplicate-link")
        for field in ("invoices", "bank_transactions", "bank_adjustments", "human_records",
                      "contracts", "fulfillments", "period_series", "accounts", "rows"):
            duplicate[field] = []
        duplicate["declarations"] = {}
        with self.assertRaisesRegex(materials.InputError, "四流勾稽编号重复"):
            docs = [primary, duplicate]
            materials.build_dataset(docs, selections(docs), COMPANY, KEYS)


class ContractLedgerWebFlow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp.name) / "contracts.db"
        self.old_store = app_module.store
        app_module.store = Store(self.db_path)
        app_module.store.create_user(
            "contract-admin", "contract-test-2026", "合同测试", "org_admin", "default"
        )
        # Isolate transient upload drafts as well as the persisted test DB.
        app = FastAPI(routes=[r for r in app_module.app.routes if not r.path.startswith("/api/materials/")],
                      exception_handlers=app_module.app.exception_handlers,
                      middleware=app_module.app.user_middleware)
        app_module.register_material_upload(app, app_module._user, app_module._allow,
                                            app_module._save_audit, app_module.RULES_DIR, app_module._audit_or_404)
        self.client = TestClient(app)
        self.client.post("/api/login", json={"username": "contract-admin", "password": "contract-test-2026"})

    def tearDown(self):
        self.client.close()
        app_module.store = self.old_store
        self.temp.cleanup()

    def test_preview_audit_and_persisted_four_flow_evidence(self):
        response = self.client.post(
            "/api/materials/preview", files=[("files", ("四流台账.xlsx", contract_workbook()))]
        )
        self.assertEqual(response.status_code, 200, response.text)
        draft = response.json()
        self.assertEqual(len(draft["documents"][0]["contracts"]), 1)
        payload = {
            "token": draft["token"], "mode": "merge", "same_scope": True,
            "company": COMPANY, "selections": list(selections(draft["documents"]).values()),
        }
        audited = self.client.post("/api/materials/audit", json=payload)
        self.assertEqual(audited.status_code, 200, audited.text)
        body = audited.json()
        self.assertFalse(body["errors"], body)
        audit_id = body["results"][0]["audit"]["audit_id"]
        saved = Store(self.db_path).get_audit(audit_id)["dataset"]
        self.assertEqual(saved.get("合同.四流完整合同数量"), 1)
        self.assertIn("四流台账.xlsx", saved.source_of("合同.四流完整勾稽金额"))

    def test_pending_and_complete_metrics_survive_history_without_changing_rule_counts(self):
        for complete in (True, False):
            with self.subTest(complete=complete):
                draft = self.client.post("/api/materials/preview", files=[("files", (
                    "归集来源.xlsx", contract_workbook(include_invoice=complete, include_bank=complete,
                                                   include_fulfillment=complete, include_link=complete),
                ))]).json()
                result = self.client.post("/api/materials/audit", json={
                    "token": draft["token"], "mode": "merge", "same_scope": True,
                    "company": COMPANY, "selections": list(selections(draft["documents"]).values()),
                })
                self.assertEqual(result.status_code, 200, result.text)
                self.assertFalse(result.json()["errors"])
                audit = result.json()["results"][0]["audit"]
                metrics = {m["name"]: m for m in audit["metrics"]}
                self.assertEqual(metrics["合同.四流完整合同数量"]["value"], "1.00" if complete else "0.00")
                self.assertEqual(metrics["合同.四流待完善合同数量"]["value"], "0.00" if complete else "1.00")
                self.assertEqual(audit["summary"]["total"], len(audit["findings"]))
                self.assertFalse(any(f["id"].startswith("合同.") for f in audit["findings"]))
                app_module.store = Store(self.db_path)
                self.assertEqual(self.client.get(f"/api/audits/{audit['audit_id']}").json(), audit)


if __name__ == "__main__":
    unittest.main()
