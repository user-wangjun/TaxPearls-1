"""Server-owned standard workbook cells and non-destructive user corrections.

Only existing numeric cells in supported statement tables can be addressed.
Rebuild those tables in memory and reuse their parsers; never mutate the retained
original or accept browser-provided sources, identifiers or derived metrics.
"""
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import re

from openpyxl import Workbook

from . import config, loader
from .input_errors import InputError

TABLES = {
    config.SHEET_ACCOUNTS: (config.COL_ACCOUNTS, (2, 3, 4, 5)),
    config.SHEET_DECLARATION: (['项目', '金额'], (1,)),
    config.SHEET_INCOME: (config.COL_STATEMENT, (1,)),
    config.SHEET_BALANCE: (config.COL_STATEMENT, (1,)),
    config.SHEET_CASHFLOW: (config.COL_STATEMENT, (1,)),
    config.SHEET_SUPPLEMENT: (config.COL_SUPPLEMENT, (1,)),
    config.SHEET_HISTORY: (config.COL_HISTORY, (1,)),
}
ACCOUNT_FIELDS = {2: 'opening', 3: 'debit', 4: 'credit', 5: 'closing'}


def capture(workbook):
    """Called only after the original standard table parsers validate headers."""
    from .materials import _serial
    return {name: [[_serial(cell) for cell in row] for row in
                   workbook[name].iter_rows(max_col=len(headers), values_only=True)]
            for name, (headers, _) in TABLES.items() if name in workbook.sheetnames}


def fields(document):
    """Original, immutable coordinates. Missing numeric values remain None."""
    output = []
    if document.get('error') or document.get('review_required'):
        return output
    for name, rows in document.get('standard_tables', {}).items():
        headers, columns = TABLES[name]
        for number, row in enumerate(rows[1:], 2):
            if row[0] is None:
                continue
            label = str(row[0]).strip()
            if name == config.SHEET_ACCOUNTS:
                label += ' ' + str(row[1] or '').strip()
            elif name in {config.SHEET_HISTORY, config.SHEET_SUPPLEMENT}:
                label += ' · ' + str(row[3] or '期间未提供')
            for column in columns:
                value = loader._number(row[column], f'{name}!{chr(65 + column)}{number}')
                output.append({'id': f'{name}!{chr(65 + column)}{number}', 'table': name,
                               'label': label + ' · ' + headers[column],
                               'value': str(value) if value is not None else None})
    return output


def normalize(document, edits):
    if not isinstance(edits, dict):
        raise InputError('标准账表修正须为单元格到数值的映射。')
    originals = {field['id']: field for field in fields(document)}
    if set(edits) - originals.keys():
        raise InputError('标准账表修正包含未知或不允许修改的单元格；请重新读取核对页。')
    normalized = {}
    for address, raw in edits.items():
        if raw is not None and (not isinstance(raw, str) or len(raw) > 100):
            raise InputError('修正值须为不超过 100 字符的数值文本或空白。')
        value = loader._number(raw, address)
        if value is not None and (abs(value) > Decimal('1e18') or value.as_tuple().exponent < -12):
            raise InputError('修正值绝对值不得超过 10^18，且最多 12 位小数。')
        original = loader._number(originals[address]['value'], address)
        if value != original:
            normalized[address] = str(value) if value is not None else None
    return normalized


def _note(address, original, edits):
    if address not in edits:
        return address
    old = '原文空白' if original is None else '原文 ' + original
    new = '留空（缺失，不是零）' if edits[address] is None else edits[address]
    label = '用户补填' if original is None else '用户修正'
    return f'{address}（{old}；{label}为 {new}；仅用于本次检测，未改原件）'


def annotate(document, original_tables, edits):
    """Propagate exact cells and corrections into computed metric sources."""
    catalogue = {f['id']: f for f in fields({'standard_tables': original_tables})}
    account_sources, declaration_sources = {}, {}
    for address, field in catalogue.items():
        sheet, cell = address.rsplit('!', 1)
        number, column = int(cell[1:]), ord(cell[0]) - 65
        row = original_tables[sheet][number - 1]
        source = _note(address, field['value'], edits)
        if sheet == config.SHEET_ACCOUNTS:
            account_sources.setdefault(str(row[0]).strip(), {})[ACCOUNT_FIELDS[column]] = source
        elif sheet == config.SHEET_DECLARATION:
            declaration_sources[str(row[0]).strip()] = source
    document['account_cell_sources'] = account_sources
    document['declaration_cell_sources'] = declaration_sources
    # Statements/supplement/period-series already carry their exact cell paths.
    for row in [*document['rows'], *document.get('period_series', [])]:
        source = row['source']
        for address in edits:
            pattern = re.escape(address) + r'(?!\d)'
            if re.search(pattern, source):
                note = _note(address, catalogue[address]['value'], edits)
                row['source'] = re.sub(pattern, lambda _: note, row['source'])
                row['detail'] = row.get('detail', '') + '；' + note


def apply(document, corrections):
    """Return candidate copy plus changes; the original document stays intact."""
    from . import materials
    edits = normalize(document, corrections)
    candidate = deepcopy(document)
    if not edits:
        return candidate, []
    tables = deepcopy(document['standard_tables'])
    for address, value in edits.items():
        sheet, cell = address.rsplit('!', 1)
        tables[sheet][int(cell[1:]) - 1][ord(cell[0]) - 65] = value
    book = Workbook()
    company = book.active
    company.title = config.SHEET_COMPANY
    company.append(['项目', '内容'])
    for label, key in zip(config.COMPANY_FIELDS, materials.COMPANY_KEYS):
        company.append([label, document['company'].get(key, '')])
    for name, rows in tables.items():
        sheet = book.create_sheet(name)
        for row in rows:
            # Source strings are literals, never formulas in the scratch copy.
            sheet.append(row)
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    output = BytesIO()
    try:
        book.save(output)
    finally:
        book.close()
    parsed = deepcopy(document)
    materials._excel(output.getvalue(), parsed, allow_incomplete_company=True, capture_standard=True)
    for key in ('accounts', 'declarations', 'rows', 'period_series'):
        candidate[key] = parsed[key]
    annotate(candidate, document['standard_tables'], edits)
    originals = {field['id']: field for field in fields(document)}
    changes = [{'file_id': document['id'], 'field': address, 'old': originals[address]['value'],
                'new': value, 'origin': 'user', 'source': document['name'] + ' / ' + address}
               for address, value in edits.items()]
    return candidate, changes
