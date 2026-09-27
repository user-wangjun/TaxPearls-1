"""Shared XLSX resource limits for every import entry point."""
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

MAX_FILE = 10 * 1024 * 1024
MAX_EXPANDED = 50 * 1024 * 1024


def open_workbook(data, *, read_only=False):
    """Check the container before openpyxl allocates workbook objects."""
    if not data or len(data) > MAX_FILE:
        raise ValueError("Excel 文件为空或超过 10MB。")
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > MAX_EXPANDED:
                raise ValueError("Excel 解压内容超过限制。")
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=read_only)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Excel 工作簿损坏或格式无效。") from exc
    if any(sheet.max_row > 30000 or sheet.max_column > 100 for sheet in workbook):
        workbook.close()
        raise ValueError("每张 Excel 工作表最多 30000 行、100 列。")
    return workbook
