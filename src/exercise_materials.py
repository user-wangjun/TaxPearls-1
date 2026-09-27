"""Application-side XLSX export of frozen synthetic teaching inputs, not answers.

Uses the same standard template and runtime Excel dependency as the loader.
No local artifact-authoring runtime is required by the deployed application.
"""
from src import periods
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from math import ceil
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, loader


def export(dataset):
    """Export only input facts; refuse any numeric round-trip loss."""
    if not dataset.company.taxpayer_id.startswith("TEST-GEN-"):
        raise ValueError("仅可导出规则生成的纯仿真教学材料。")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "TaxPearls"
    wb.properties.title = "纯仿真教学材料（不含答案）"
    wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)

    def sheet(name, headers, rows, widths):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = "1:1"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.oddHeader.center.text = "纯仿真教学材料"
        ws.oddFooter.center.text = "金额单位元；比例为小数；人数为整数"
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in ws:
            lines = 1
            for cell in row:
                numeric = isinstance(cell.value, (int, float))
                cell.font = Font(name="Arial", size=11, color="172B4D")
                cell.alignment = Alignment(horizontal="right" if numeric else "left", vertical="center", wrap_text=True)
                if numeric:
                    cell.number_format = "#,##0.00;[Red](#,##0.00)"
                elif isinstance(cell.value, str):
                    # Even source text starting '=' remains literal, never a formula.
                    cell.data_type = "s"
                    lines = max(lines, ceil(len(cell.value) * 2 / widths[cell.column-1]))
            ws.row_dimensions[row[0].row].height = max(24, lines * 17)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="24476B")
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        return ws

    company = dataset.company
    sheet(config.SHEET_COMPANY, ["项目", "内容"],
          list(zip(config.COMPANY_FIELDS, (company.name, company.taxpayer_id, company.industry, company.period)))
          + [("材料说明", "纯仿真标准化账套及底稿。不是真实票据或完整凭证账簿。不含标准答案。")], [24, 80])
    sheet(config.SHEET_ACCOUNTS, config.COL_ACCOUNTS,
          [(a.code, a.name, a.opening, a.debit, a.credit, a.closing) for a in dataset.accounts],
          [20, 32, 22, 22, 22, 22])
    sheet(config.SHEET_DECLARATION, ["项目", "金额"], sorted(dataset.declarations.items()), [32, 24])
    mapped = set(loader._build_metrics(dataset.accounts, dataset.declarations))
    for name, prefix in ((config.SHEET_INCOME, "利润表."), (config.SHEET_BALANCE, "资产负债表."),
                         (config.SHEET_CASHFLOW, "现金流量表.")):
        keys = sorted(k for k in dataset.metrics if k.startswith(prefix))
        sheet(name, config.COL_STATEMENT, [(k[len(prefix):], dataset.metrics[k].value) for k in keys], [40, 24])
        mapped.update(keys)
    historical, supplement = [], []
    year = periods.parse_period(company.period, "教学年度").end.year
    for key, metric in sorted(dataset.metrics.items()):
        if key in mapped:
            continue
        period = company.period
        if key.startswith(("历史.", "年度.")):
            actual = year - (2 if ".前年" in key else (1 if key.startswith("历史.") or ".上年" in key else 0))
            period = str(actual)
            rows = historical
        else:
            rows = supplement
        rows.append((key, metric.value, metric.source, period, metric.detail))
    for name, headers, rows in ((config.SHEET_HISTORY, config.COL_HISTORY, historical),
                               (config.SHEET_SUPPLEMENT, config.COL_SUPPLEMENT, supplement)):
        ws = sheet(name, headers, rows, [38, 24, 48, 16, 100])
        for idx, row in enumerate(rows, 2):
            if row[0].endswith("人数"):
                ws.cell(idx, 2).number_format = "#,##0"
            elif row[0].startswith("参考.") or row[0].endswith("税率"):
                ws.cell(idx, 2).number_format = "0.00%"
    raw = BytesIO()
    wb.save(raw)
    wb.close()
    # Fix OOXML clock metadata and ZIP timestamps so repeat downloads are stable.
    stable = BytesIO()
    with ZipFile(raw) as source, ZipFile(stable, "w", compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                root = ElementTree.fromstring(content)
                root.find("{http://purl.org/dc/terms/}modified").text = "2000-01-01T00:00:00Z"
                content = ElementTree.tostring(root, encoding="utf-8")
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, content)
    result = stable.getvalue()
    restored = loader.load_bytes(result)
    if (restored.company != company or restored.accounts != dataset.accounts
            or restored.declarations != dataset.declarations or restored.values() != dataset.values()):
        raise ValueError("导出材料重新导入后数值不一致，未提供下载。")
    return result
